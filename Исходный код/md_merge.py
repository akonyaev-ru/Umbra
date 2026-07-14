"""Трёхстороннее слияние: md-ответ ИИ -> правки в анонимизированную копию оригинала.

Вход: текст ответа ИИ, base_blocks из md_serializer.serialize_document (той же
повторной сериализации оригинала) и живой Document. Md — только транспорт diff'а:
итоговый docx никогда не строится из md, правится сам Document.

Гарантии:
- блок, совпавший с базой (equal), НЕ ТРОГАЕТСЯ — форматирование байт-в-байт;
- правленый абзац меняется пословным diff'ом: не тронутые ИИ слова остаются в
  своих runs (сохраняют начертание);
- структура таблиц всегда из оригинала; правки в ячейках — только при полном
  совпадении размерности;
- удаления НЕ применяются молча: собираются в deletions и ждут подтверждения
  (кроме «слияния абзацев», где текст сохранён в первом абзаце пары).

Выход: MergeReport (см. merge_ai_md) + список отложенных удалений.
"""

import re
import difflib
from copy import deepcopy

from markdown_it import MarkdownIt
from docx.oxml.ns import qn

from md_serializer import (protect_labels, restore_labels, normalize_for_align,
                           IMG_TOKEN_RE, FOOTREF_RE, _ENUM_PREFIX_RE, LABEL_RE_CI)

# Пороги выравнивания (калибруются на корпусе; до калибровки — режим «бета»).
SIM_THRESHOLD = 0.55        # минимальное сходство 1:1, чтобы считать «правкой»
JOIN_THRESHOLD = 0.80       # сходство для склейки/разбиения 1:2, 2:1
DELETED_CLUSTER_MIN = 3     # подряд удалённых блоков => подозрение на усечение
COVERAGE_MIN = 0.70         # ниже — «ИИ вернул неполный документ»

_NBSP = chr(0xA0)
# Нулевой ширины: word-joiner U+2060 (мы сами добавляем в md), zero-width space,
# BOM. Срезаем в ответе ИИ, чтобы невидимый символ не запёкся в итоговый docx и
# не сбивал сравнение блоков. Строим через chr() — в исходнике нет невидимых.
_ZW_RE = re.compile('[' + chr(0x2060) + chr(0x200B) + chr(0xFEFF) + ']')
_FOOTDEF_RE = re.compile(r'^\[\^([fe])(\d{1,4})\]:\s*(.*)$', re.DOTALL)
_FENCE_RE = re.compile(r'^```[a-zA-Z]*\s*\n(.*)\n```\s*$', re.DOTALL)
_CHATTER_RE = re.compile(
    r'(?:вот|ниже|исправленн|обновлённ|обновленн|итогов|переработанн|готов)', re.IGNORECASE)

# Фразы из нашей шапки-инструкции ([ANON].md): если ИИ отдал их эхом как текст,
# такие блоки отбрасываются на уровне парсинга — второй рубеж после preclean.
# Иначе они «вставятся» в документ и деанон подставит в легенду реальные ПДн.
_INSTRUCTION_MARKERS = (
    'Инструкция для ИИ',
    'персональные данные заменены метками',
    'не изменяйте, не переводите и не удаляйте метки',
    'структуру не менять».',
    'оставляйте на месте.',
    'Верните документ ЦЕЛИКОМ',
    'В документе использованы метки',
)


def _is_instruction_echo(text):
    return any(m in text for m in _INSTRUCTION_MARKERS)


# --------------------------------------------------------------------------- #
#  Предочистка ответа ИИ                                                       #
# --------------------------------------------------------------------------- #

def _strip_leading_quote_instruction(lines):
    """Срезает ведущий blockquote с эхом нашей инструкции."""
    i = 0
    while i < len(lines) and not lines[i].strip():
        i += 1
    if i < len(lines) and lines[i].lstrip().startswith('>'):
        j = i
        quote = []
        while j < len(lines) and (lines[j].lstrip().startswith('>') or not lines[j].strip()):
            quote.append(lines[j])
            j += 1
            if j < len(lines) and not lines[j].strip() and (j + 1 >= len(lines) or not lines[j + 1].lstrip().startswith('>')):
                break
        qtext = '\n'.join(quote)
        if 'нструкция' in qtext or 'метк' in qtext.lower():
            return lines[:i] + lines[j:], True
    return lines, False


def _strip_leading_chatter(lines):
    """Срезает вступительную реплику ассистента («Вот исправленный документ:»)."""
    i = 0
    while i < len(lines) and not lines[i].strip():
        i += 1
    if i < len(lines):
        first = lines[i].strip()
        if (len(first) < 140 and first.endswith(':') and _CHATTER_RE.search(first)
                and '[' not in first and not first.startswith('#')):
            return lines[:i] + lines[i + 1:], True
    return lines, False


def preclean(text):
    """Срезает обёртки чат-ответа: BOM, цельный код-фенс, эхо инструкции,
    вступительную реплику ассистента. Реплика может стоять ПЕРЕД эхом
    инструкции — проходы повторяются до стабилизации."""
    text = text.lstrip('﻿').replace('\r\n', '\n').replace('\r', '\n')
    text = _ZW_RE.sub('', text)

    m = _FENCE_RE.match(text.strip())
    if m:
        text = m.group(1)

    lines = text.split('\n')
    for _ in range(4):
        lines, c1 = _strip_leading_chatter(lines)
        lines, c2 = _strip_leading_quote_instruction(lines)
        if not (c1 or c2):
            break

    return '\n'.join(lines)


# --------------------------------------------------------------------------- #
#  Разбор md от ИИ в блоки                                                     #
# --------------------------------------------------------------------------- #

def _inline_text(token):
    """Плоский текст inline-токена: содержимое без разметки (начертание правленых
    фрагментов берётся из оригинала/соседей, а не из ** ** ответа ИИ)."""
    if token is None:
        return ''
    out = []
    for ch in (token.children or []):
        t = ch.type
        if t in ('text', 'code_inline'):
            out.append(ch.content)
        elif t in ('softbreak', 'hardbreak'):
            out.append(' ')
        elif t == 'image':
            out.append(ch.content)
        # emphasis/strong/s открывающие-закрывающие токены пропускаем: их
        # содержимое приходит вложенными text-токенами.
        elif ch.children:
            out.append(_inline_text(ch))
    return ''.join(out)


def _html_table_cells(html):
    """Тексты ячеек из html-блока (замороженная таблица)."""
    cells = re.findall(r'<t[dh][^>]*>(.*?)</t[dh]>', html, re.DOTALL | re.IGNORECASE)
    return [re.sub(r'<[^>]+>', ' ', c).strip() for c in cells]


def parse_ai_blocks(text, known_labels):
    """Разбирает ответ ИИ в блоки [{'kind','text','norm','heading','cells'}].
    Метки защищаются PUA до парсинга и восстанавливаются в канонической форме."""
    prot, registry = protect_labels(text, known=known_labels)

    md = MarkdownIt('commonmark').enable('table').enable('strikethrough')
    # reference definitions и indented code выключены: «[МЕТКА]: …» и абзацы с
    # отступом — это текст документа, а не синтаксис.
    try:
        md.disable(['reference'])
    except Exception:
        pass
    try:
        md.disable(['code'])
    except Exception:
        pass

    tokens = md.parse(prot)
    blocks = []

    def add_para(txt, heading=None):
        txt = restore_labels(txt, registry).strip()
        if not txt or _is_instruction_echo(txt):
            return
        m = _FOOTDEF_RE.match(txt)
        if m:
            blocks.append({'kind': 'footnote', 'note_key': (m.group(1), m.group(2)),
                           'text': m.group(3).strip(),
                           'norm': normalize_for_align(m.group(3))})
            return
        blocks.append({'kind': 'para', 'text': txt,
                       'norm': normalize_for_align(txt), 'heading': heading})

    i = 0
    while i < len(tokens):
        tok = tokens[i]
        t = tok.type
        if t == 'heading_open':
            level = int(tok.tag[1]) if tok.tag and tok.tag[1:].isdigit() else 1
            add_para(_inline_text(tokens[i + 1]), heading=level)
            i += 3
        elif t == 'paragraph_open':
            add_para(_inline_text(tokens[i + 1]))
            i += 3
        elif t == 'inline':
            add_para(_inline_text(tok))
            i += 1
        elif t == 'fence' or t == 'code_block':
            for line in tok.content.split('\n'):
                add_para(line)
            i += 1
        elif t == 'html_block':
            content = restore_labels(tok.content, registry)
            if '<table' in content.lower():
                cells = [restore_labels(c, registry) for c in _html_table_cells(tok.content)]
                blocks.append({'kind': 'frozen', 'text': content,
                               'norm': normalize_for_align(' '.join(cells)),
                               'cells': cells})
            else:
                add_para(re.sub(r'<[^>]+>', ' ', content))
            i += 1
        elif t == 'table_open':
            rows, row = [], None
            j = i
            while j < len(tokens) and tokens[j].type != 'table_close':
                tt = tokens[j].type
                if tt == 'tr_open':
                    row = []
                elif tt == 'tr_close':
                    rows.append(row)
                    row = None
                elif tt == 'inline' and row is not None:
                    row.append(restore_labels(_inline_text(tokens[j]), registry).strip())
                j += 1
            flat = [c for r in rows for c in r]
            blocks.append({'kind': 'table', 'rows': rows,
                           'text': ' '.join(flat),
                           'norm': normalize_for_align(' '.join(flat))})
            i = j + 1
        else:
            i += 1
    return blocks


# --------------------------------------------------------------------------- #
#  Выравнивание                                                                #
# --------------------------------------------------------------------------- #

_RATIO_LEN_CAP = 4000        # выше — грубая оценка вместо точного O(L^2) ratio


def _ratio(a, b):
    la, lb = len(a), len(b)
    if not la or not lb:
        return 1.0 if la == lb else 0.0
    sm = difflib.SequenceMatcher(None, a, b, autojunk=False)
    # real_quick_ratio (O(1) по длинам) и quick_ratio (O(L) по мультимножеству
    # символов) — ВЕРХНИЕ границы ratio(). Если они ниже порога сходства, точный
    # O(L^2) ratio() не нужен — так снимается квадратичное зависание на больших /
    # полностью переписанных документах и на гигантских таблицах (один _ratio на
    # двух огромных нормах).
    if sm.real_quick_ratio() < SIM_THRESHOLD or sm.quick_ratio() < SIM_THRESHOLD:
        return 0.0
    if la > _RATIO_LEN_CAP or lb > _RATIO_LEN_CAP:
        return sm.quick_ratio()
    return sm.ratio()


def _labels(text):
    """Мультимножество канонических меток [КАТ_N] в тексте (отсортировано).
    Инвариант слияния: набор меток блока при правке ИИ меняться не должен — иначе
    ИИ подменил номер (данные другого лица), потерял метку (потеря значения) или
    добавил чужую; правку в таком блоке применять НЕЛЬЗЯ (см. modify_para)."""
    return sorted('[' + m.group(1).upper() + '_' + m.group(2) + ']'
                  for m in LABEL_RE_CI.finditer(text))


def _dp_align(base_win, ai_win):
    """Локальное выравнивание внутри replace-окна: DP по сумме сходств с платой
    за пропуск. Возвращает список пар (bi, ai) в порядке следования."""
    n, m = len(base_win), len(ai_win)
    GAP = 0.20
    score = [[0.0] * (m + 1) for _ in range(n + 1)]
    back = [[None] * (m + 1) for _ in range(n + 1)]
    for i in range(1, n + 1):
        score[i][0] = score[i - 1][0] - GAP
        back[i][0] = 'up'
    for j in range(1, m + 1):
        score[0][j] = score[0][j - 1] - GAP
        back[0][j] = 'left'
    for i in range(1, n + 1):
        for j in range(1, m + 1):
            r = _ratio(base_win[i - 1], ai_win[j - 1])
            diag = score[i - 1][j - 1] + (r if r >= SIM_THRESHOLD else -1.0)
            up = score[i - 1][j] - GAP
            left = score[i][j - 1] - GAP
            best = max(diag, up, left)
            score[i][j] = best
            back[i][j] = 'diag' if best == diag else ('up' if best == up else 'left')
    pairs = []
    i, j = n, m
    while i > 0 or j > 0:
        d = back[i][j]
        if d == 'diag':
            i, j = i - 1, j - 1
            if _ratio(base_win[i], ai_win[j]) >= SIM_THRESHOLD:
                pairs.append((i, j))
        elif d == 'up':
            i -= 1
        else:
            j -= 1
    pairs.reverse()
    return pairs


def align_blocks(base_blocks, ai_blocks):
    """Выравнивает базовые блоки с блоками ответа ИИ.

    Возвращает список операций:
      ('equal', bi, ai) ('modified', bi, ai) ('merged', (bi1, bi2), ai)
      ('split', bi, (ai1, ai2)) ('deleted', bi) ('inserted', ai, prev_bi)
    prev_bi — индекс последнего базового блока перед вставкой (или -1).
    Сноски в выравнивании не участвуют (сопоставляются по ключу отдельно)."""
    base_norms = [b['norm'] for b in base_blocks]
    ai_norms = [a['norm'] for a in ai_blocks]

    sm = difflib.SequenceMatcher(None, base_norms, ai_norms, autojunk=False)
    ops = []
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == 'equal':
            for k in range(i2 - i1):
                ops.append(('equal', i1 + k, j1 + k))
            continue
        b_idx = list(range(i1, i2))
        a_idx = list(range(j1, j2))
        if tag == 'delete':
            ops.extend(('deleted', bi) for bi in b_idx)
            continue
        if tag == 'insert':
            prev = i1 - 1
            ops.extend(('inserted', ai, prev) for ai in a_idx)
            continue

        # replace: локальное DP-выравнивание.
        pairs = _dp_align([base_norms[i] for i in b_idx], [ai_norms[j] for j in a_idx])
        matched_b = {b_idx[p[0]] for p in pairs}
        matched_a = {a_idx[p[1]] for p in pairs}
        for pb, pa in pairs:
            bi, ai = b_idx[pb], a_idx[pa]
            if base_norms[bi] == ai_norms[ai]:
                ops.append(('equal', bi, ai))
            else:
                ops.append(('modified', bi, ai))

        rest_b = [i for i in b_idx if i not in matched_b]
        rest_a = [j for j in a_idx if j not in matched_a]

        # Склейка 2:1 и разбиение 1:2 на остатке.
        used_b, used_a = set(), set()
        for bi in list(rest_b):
            if bi in used_b or bi + 1 not in rest_b:
                continue
            for aj in rest_a:
                if aj in used_a:
                    continue
                joined = base_norms[bi] + ' ' + base_norms[bi + 1]
                if _ratio(joined, ai_norms[aj]) >= JOIN_THRESHOLD:
                    ops.append(('merged', (bi, bi + 1), aj))
                    used_b.update((bi, bi + 1))
                    used_a.add(aj)
                    break
        for bi in list(rest_b):
            if bi in used_b:
                continue
            for aj in rest_a:
                if aj in used_a or aj + 1 not in rest_a or aj + 1 in used_a:
                    continue
                joined = ai_norms[aj] + ' ' + ai_norms[aj + 1]
                if _ratio(base_norms[bi], joined) >= JOIN_THRESHOLD:
                    ops.append(('split', bi, (aj, aj + 1)))
                    used_b.add(bi)
                    used_a.update((aj, aj + 1))
                    break

        for bi in rest_b:
            if bi not in used_b:
                ops.append(('deleted', bi))
        prev = b_idx[0] - 1
        for aj in rest_a:
            if aj not in used_a:
                ops.append(('inserted', aj, prev))

    # Пост-проход: DP мог сматчить «первый абзац ↔ объединённый текст» как
    # modified, оставив второй абзац в deleted. Если пара (bi, bi+1) вместе
    # похожа на ai-блок сильнее порога склейки — это merged (второй абзац
    # удаляется автоматически: его текст сохранён в первом).
    by_key = {(op[0], op[1]): op for op in ops}
    for op in list(ops):
        if op[0] != 'modified':
            continue
        bi, ai = op[1], op[2]
        dele = by_key.get(('deleted', bi + 1))
        if dele is not None and dele in ops:
            joined = base_norms[bi] + ' ' + base_norms[bi + 1]
            if _ratio(joined, ai_norms[ai]) >= JOIN_THRESHOLD:
                ops.remove(op)
                ops.remove(dele)
                ops.append(('merged', (bi, bi + 1), ai))
                continue
        # Симметрично: ИИ разбил абзац — modified + inserted сразу после него.
        for ins in ops:
            if ins[0] == 'inserted' and ins[2] == bi and ins[1] == ai + 1:
                joined = ai_norms[ai] + ' ' + ai_norms[ai + 1]
                if _ratio(base_norms[bi], joined) >= JOIN_THRESHOLD:
                    ops.remove(op)
                    ops.remove(ins)
                    ops.append(('split', bi, (ai, ai + 1)))
                break

    # Перестановки: удалённый и вставленный блоки с одинаковой нормой — это move.
    deleted = {op[1]: op for op in ops if op[0] == 'deleted'}
    inserted = [op for op in ops if op[0] == 'inserted']
    moves = []
    for ins in inserted:
        ai = ins[1]
        for bi, dop in list(deleted.items()):
            if base_norms[bi] == ai_norms[ai]:
                moves.append((bi, ai))
                ops.remove(dop)
                ops.remove(ins)
                ops.append(('moved', bi, ai))
                del deleted[bi]
                break
    return ops


# --------------------------------------------------------------------------- #
#  Применение правок                                                           #
# --------------------------------------------------------------------------- #

def _prepare_target(ai_text, base_meta):
    """Готовит текст ИИ к записи в docx: срезает служебные токены и литеральную
    нумерацию (видимый номер даёт numPr Word, а не текст), возвращает NBSP в
    разряды чисел (equal-абзацы это получают бесплатно из оригинала)."""
    t = _ZW_RE.sub('', ai_text)
    t = FOOTREF_RE.sub('', t)
    t = IMG_TOKEN_RE.sub('', t)
    if base_meta.get('num_literal'):
        t = _ENUM_PREFIX_RE.sub('', t, count=1)
        if t.startswith('- '):
            t = t[2:]
    # «100 000» -> «100<NBSP>000» (после правок ИИ обычный пробел в числах).
    prev = None
    while prev != t:
        prev = t
        t = re.sub(r'(?<=\d) (?=\d{3}(?:\D|$))', _NBSP, t)
    return re.sub(r'\s+', ' ', t).strip()


def _word_diff_spans(old, new):
    """Пословный diff old->new как список замен для _apply_span_replacements
    (справа-налево). Не тронутые слова остаются в своих runs.

    Вставки (i1==i2) _apply_span_replacements не берёт (нулевой диапазон не
    перекрывает ни один run), поэтому расширяем их на один соседний символ —
    он в equal-зоне (difflib разделяет не-equal opcodes равными), значит
    переписать его тем же значением безопасно."""
    sm = difflib.SequenceMatcher(None, old, new, autojunk=False)
    reps = []
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == 'equal':
            continue
        ins = new[j1:j2]
        if i1 == i2:                       # чистая вставка
            if i1 > 0:
                i1 -= 1
                ins = old[i1] + ins
            elif i2 < len(old):
                i2 += 1
                ins = ins + old[i2 - 1]
            elif not old:
                # Пустой абзац: некуда расширяться — вставку выполнит вызывающий
                # (запись в единственный run напрямую).
                pass
        reps.append({'start': i1, 'stop': i2, 'ins': ins})
    reps.sort(key=lambda r: r['start'], reverse=True)
    return reps


def _p_has_numpr(p_el):
    """True, если у абзаца-элемента есть прямой numPr: Word сам нарисует номер,
    поэтому литеральную нумерацию из текста ИИ надо срезать (иначе «10. 10. …»)."""
    if p_el is None:
        return False
    pPr = p_el.find(qn('w:pPr'))
    return pPr is not None and pPr.find(qn('w:numPr')) is not None


def _clone_paragraph_after(anchor_el, text, template_p, before=False):
    """Вставляет новый абзац с текстом рядом с anchor_el (после; либо ПЕРЕД при
    before=True — для вставки в самое начало документа), клонируя pPr (стиль,
    numPr, отступы) и rPr первого run из template_p."""
    new_p = deepcopy(template_p)
    # Удаляем всё содержимое, кроме pPr.
    for child in list(new_p):
        if child.tag != qn('w:pPr'):
            new_p.remove(child)
    # Один run: rPr из первого run шаблона (если был).
    r_el = None
    for r in template_p.findall(qn('w:r')):
        r_el = r
        break
    run = deepcopy(r_el) if r_el is not None else None
    if run is None:
        from docx.oxml import OxmlElement
        run = OxmlElement('w:r')
    else:
        for child in list(run):
            if child.tag != qn('w:rPr'):
                run.remove(child)
    from docx.oxml import OxmlElement
    t_el = OxmlElement('w:t')
    t_el.text = text
    t_el.set(qn('xml:space'), 'preserve')
    run.append(t_el)
    new_p.append(run)
    if before:
        anchor_el.addprevious(new_p)
    else:
        anchor_el.addnext(new_p)
    return new_p


def _block_anchor_el(block):
    ref = block['ref']
    if hasattr(ref, '_p'):
        return ref._p
    if hasattr(ref, '_tbl'):
        return ref._tbl
    return None


def merge_ai_md(ai_text, base_blocks, mapping, notes, apply_spans):
    """Главная точка: применяет ответ ИИ к документу (через base_blocks[i]['ref']).

    apply_spans(paragraph, reps) — функция пофрагментной замены (из
    DocumentProcessor._apply_span_replacements), переиспользуем её сохранение
    форматирования.

    Возвращает (report, deletions):
    report = {'coverage','equal','modified','inserted','deleted','moved',
              'split_joined','merged_pairs','table_conflicts','frozen_touched',
              'format_heuristics','truncation_suspect','deleted_clusters',
              'notes_modified','notes_missing','warnings'}
    deletions = [{'bi','preview','ref','auto'}] — физически НЕ удалены."""
    known = set(mapping)
    text = preclean(ai_text)
    ai_all = parse_ai_blocks(text, known)

    # Сноски: сопоставление по ключу, вне общего выравнивания.
    ai_notes = {b['note_key']: b for b in ai_all if b['kind'] == 'footnote'}
    ai_blocks = [b for b in ai_all if b['kind'] != 'footnote']
    base_notes = [b for b in base_blocks if b['kind'] == 'footnote']
    base_main = [b for b in base_blocks if b['kind'] != 'footnote']

    report = {'coverage': 0.0, 'equal': 0, 'modified': 0, 'inserted': 0,
              'deleted': 0, 'moved': 0, 'split_joined': 0, 'merged_pairs': 0,
              'table_conflicts': [], 'frozen_touched': [], 'format_heuristics': [],
              'truncation_suspect': False, 'deleted_clusters': [],
              'notes_modified': 0, 'notes_missing': [], 'warnings': [],
              'changed_texts': [], 'label_mismatch': 0, 'merged_kept': 0}
    deletions = []

    ops = align_blocks(base_main, ai_blocks)

    def modify_para(block, target_text, expected_labels=None):
        para = block['ref']
        old = para.text
        target = _prepare_target(target_text, block.get('meta', {}))
        if normalize_for_align(old) == normalize_for_align(target):
            return False
        # ИНВАРИАНТ СОХРАННОСТИ МЕТОК: если набор меток в тексте ИИ отличается от
        # ожидаемого — метку переименовали/сменили номер ([ФИО_1]→[ФИО_2] = данные
        # ДРУГОГО лица) или потеряли (значение исчезло бы молча) — правку НЕ
        # применяем, оставляем оригинал (в нём метки и значения верны). Теряем
        # только переформулировку блока ИИ — безопасный размен. expected_labels:
        # для обычной правки берём метки оригинала, для склейки — объединённые.
        exp = _labels(old) if expected_labels is None else expected_labels
        if _labels(target) != exp:
            report['label_mismatch'] += 1
            return False
        # Абзац без runs (пустой): спаны класть некуда — пишем текст напрямую
        # (форматирования у пустого абзаца всё равно нет).
        if not para.runs:
            if not target:
                return False
            para.text = target
            report['changed_texts'].append(target)
            return True
        reps = _word_diff_spans(old, target)
        if reps:
            apply_spans(para, reps)
            report['changed_texts'].append(target)
            # Если правка легла на стык runs с разным начертанием — формат выбран
            # по левому соседу; такие места идут в отчёт на глазную проверку.
            if len({(r.bold, r.italic) for r in para.runs if r.text.strip()}) > 1:
                report['format_heuristics'].append(target[:90])
        return bool(reps)

    matched_base = 0
    # Последовательные вставки с одним prev должны идти В ПОРЯДКЕ ответа ИИ:
    # якорь каждой следующей — элемент, вставленный предыдущей (иначе addnext
    # к одному якорю развернёт их задом наперёд).
    last_inserted_after = {}
    for op in ops:
        kind = op[0]
        if kind == 'equal':
            report['equal'] += 1
            matched_base += 1
        elif kind == 'moved':
            report['moved'] += 1
            matched_base += 1
        elif kind == 'modified':
            b, a = base_main[op[1]], ai_blocks[op[2]]
            matched_base += 1
            if b['kind'] == 'para':
                if modify_para(b, a['text']):
                    report['modified'] += 1
                else:
                    report['equal'] += 1
            elif b['kind'] == 'table':
                if a.get('rows') and b['meta']['dims'] == (len(a['rows']), len(a['rows'][0]) if a['rows'] else 0):
                    table = b['ref']
                    changed = False
                    for ri, row in enumerate(table.rows):
                        for ci, cell in enumerate(row.cells):
                            new_txt = a['rows'][ri][ci] if ci < len(a['rows'][ri]) else ''
                            old_txt = cell.text
                            if normalize_for_align(old_txt) == normalize_for_align(new_txt):
                                continue
                            # Метки ячейки при правке меняться не должны: если ИИ
                            # потерял «|», markdown-it добьёт строку пустыми
                            # ячейками — проверка размерности пройдёт, но значения
                            # сдвинутся и метки подменятся/потеряются. Расхождение
                            # меток → оставляем ячейку оригинала.
                            if _labels(old_txt) != _labels(new_txt):
                                report['table_conflicts'].append(
                                    'ячейка с расхождением меток — оставлена оригинальная: ' + old_txt[:40])
                                continue
                            cell_para = cell.paragraphs[0]
                            new_plain = new_txt.replace('<br>', ' ')
                            if len(cell.paragraphs) == 1 and not cell_para.runs:
                                # Пустая ячейка: спаны некуда класть.
                                cell_para.text = new_plain
                                report['changed_texts'].append(new_txt)
                                changed = True
                                continue
                            full_old = cell.text
                            reps = _word_diff_spans(full_old, new_plain)
                            if len(cell.paragraphs) == 1 and reps:
                                apply_spans(cell_para, reps)
                                report['changed_texts'].append(new_txt)
                                changed = True
                            else:
                                report['table_conflicts'].append(
                                    'многоабзацная ячейка — правка ИИ не применена: ' + new_txt[:60])
                    if changed:
                        report['modified'] += 1
                else:
                    report['table_conflicts'].append(
                        'структура таблицы изменена ИИ — оставлена оригинальная (' + b['norm'][:60] + '…)')
            elif b['kind'] == 'frozen':
                report['frozen_touched'].append(
                    'замороженная таблица изменена ИИ — оставлена оригинальная (' + b['norm'][:60] + '…)')
                matched_base += 0
        elif kind == 'merged':
            (b1, b2), a = op[1], op[2]
            blk1, blk2 = base_main[b1], base_main[b2]
            matched_base += 2
            if blk1['kind'] == 'para' and blk2['kind'] == 'para':
                # Склейка безопасна, только если текст ИИ несёт ВСЕ метки обоих
                # абзацев: тогда пишем его в первый и удаляем второй (его данные
                # сохранены в первом). Иначе НЕ сливаем и НЕ удаляем — оставляем
                # оба абзаца оригинала, иначе метка/значение второго исчезли бы
                # молча (автоудаление раньше применялось без подтверждения).
                combined = sorted(_labels(blk1['ref'].text) + _labels(blk2['ref'].text))
                if modify_para(blk1, ai_blocks[a]['text'], expected_labels=combined):
                    deletions.append({'bi': b2, 'preview': blk2['norm'][:90],
                                      'ref': blk2['ref'], 'auto': True})
                    report['merged_pairs'] += 1
                else:
                    report['merged_kept'] += 1
        elif kind == 'split':
            b, (a1, a2) = op[1], op[2]
            blk = base_main[b]
            matched_base += 1
            if blk['kind'] == 'para':
                joined = ai_blocks[a1]['text'] + ' ' + ai_blocks[a2]['text']
                if modify_para(blk, joined):
                    report['split_joined'] += 1
        elif kind == 'deleted':
            blk = base_main[op[1]]
            deletions.append({'bi': op[1], 'preview': blk['norm'][:90],
                              'ref': blk['ref'], 'auto': False})
            report['deleted'] += 1
        elif kind == 'inserted':
            a, prev = op[1], op[2]
            ai_b = ai_blocks[a]
            if ai_b['kind'] != 'para':
                report['warnings'].append('новая таблица от ИИ не переносится автоматически: '
                                          + ai_b['norm'][:60])
                continue
            # Якорь: элемент предыдущей вставки (цепочка) либо ближайший
            # предыдущий базовый блок с элементом в теле.
            anchor_el, template_p = last_inserted_after.get(prev), None
            insert_before = False
            if anchor_el is None:
                for bi in range(prev, -1, -1):
                    anchor_el = _block_anchor_el(base_main[bi])
                    if anchor_el is not None:
                        template_p = anchor_el if anchor_el.tag == qn('w:p') else None
                        break
            if anchor_el is None:
                first_el = next((_block_anchor_el(b) for b in base_main
                                 if _block_anchor_el(b) is not None), None)
                if first_el is None:
                    report['warnings'].append('не найден якорь для вставки: ' + ai_b['norm'][:60])
                    continue
                # prev < 0 → вставка ПЕРЕД первым блоком (иначе новый заголовок/
                # преамбула уехал бы за первый оригинальный абзац).
                insert_before = True
                anchor_el, template_p = first_el, first_el if first_el.tag == qn('w:p') else None
            # Шаблон стиля: для заголовков — базовый заголовок того же уровня.
            if ai_b.get('heading'):
                for b2 in base_main:
                    if b2['kind'] == 'para' and b2['meta'].get('heading') == ai_b['heading']:
                        template_p = b2['ref']._p
                        break
            if template_p is None:
                for b2 in base_main:
                    if b2['kind'] == 'para':
                        template_p = b2['ref']._p
                        break
            if template_p is None:
                report['warnings'].append('нет абзаца-шаблона для вставки: ' + ai_b['norm'][:60])
                continue
            # У шаблона есть numPr → Word сам проставит номер: срезаем литеральную
            # нумерацию из текста ИИ, иначе получится «10. 10. …».
            ins_meta = {'num_literal': True} if _p_has_numpr(template_p) else {}
            target = _prepare_target(ai_b['text'], ins_meta)
            new_el = _clone_paragraph_after(anchor_el, target, template_p, before=insert_before)
            last_inserted_after[prev] = new_el
            report['changed_texts'].append(target)
            report['inserted'] += 1

    # Сноски.
    for nb in base_notes:
        key = nb['ref']
        ab = ai_notes.get(key)
        if ab is None:
            report['notes_missing'].append('[^%s%s]' % key)
            continue
        if nb['norm'] == ab['norm']:
            continue
        paras = notes.by_id.get(key) if notes is not None else None
        if not paras:
            continue
        para = paras[0]
        target = _prepare_target(ab['text'], {})
        # Тот же инвариант меток, что и для абзацев (см. modify_para): при
        # расхождении меток оставляем оригинал сноски.
        if _labels(para.text) != _labels(target):
            report['label_mismatch'] += 1
            continue
        reps = _word_diff_spans(para.text, target)
        if reps:
            apply_spans(para, reps)
            notes.mark_changed(*key)
            report['notes_modified'] += 1
            report['changed_texts'].append(target)

    # Покрытие и усечение.
    total_base = len(base_main)
    report['coverage'] = matched_base / total_base if total_base else 1.0

    del_bis = sorted(d['bi'] for d in deletions if not d['auto'])
    cluster = []
    for bi in del_bis:
        if cluster and bi == cluster[-1] + 1:
            cluster.append(bi)
        else:
            if len(cluster) >= DELETED_CLUSTER_MIN:
                report['deleted_clusters'].append(list(cluster))
            cluster = [bi]
    if len(cluster) >= DELETED_CLUSTER_MIN:
        report['deleted_clusters'].append(list(cluster))

    if report['coverage'] < COVERAGE_MIN or report['deleted_clusters']:
        report['truncation_suspect'] = True

    return report, deletions


def apply_deletions(deletions, confirmed=None):
    """Физически удаляет подтверждённые блоки. confirmed — множество индексов в
    списке deletions; None = только автоматические (слияние абзацев)."""
    removed = 0
    for i, d in enumerate(deletions):
        if not d['auto'] and (confirmed is None or i not in confirmed):
            continue
        el = _block_anchor_el({'ref': d['ref']})
        if el is not None and el.getparent() is not None:
            el.getparent().remove(el)
            removed += 1
    return removed
