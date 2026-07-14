"""Копирование результата анонимизации в буфер обмена сразу в ДВУХ форматах:
файл (CF_HDROP) и его текст (CF_UNICODETEXT). Целевое приложение само берёт
нужный формат: область загрузки файлов / письмо / папка — получит файл, поле
чата ИИ — вставит текст. Один клик покрывает оба сценария отправки в ИИ.

Только Windows (pywin32 уже в зависимостях — используется в file_detector).
"""

import os
import struct
import win32clipboard


def extract_text(path):
    """Текст результата для вставки в поле чата ИИ. .md/.txt — как есть;
    .docx — тело + таблицы (форматирование в буфере всё равно не нужно)."""
    ext = os.path.splitext(path)[1].lower()
    if ext in ('.md', '.txt'):
        with open(path, 'r', encoding='utf-8') as f:
            return f.read()
    if ext == '.docx':
        from docx import Document
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        return '\n'.join(p for p in parts if p.strip())
    return ''


def _hdrop_blob(paths):
    """DROPFILES-структура для CF_HDROP: заголовок 20 байт + список путей в
    UTF-16-LE с двойным нулём на конце."""
    # DROPFILES { DWORD pFiles=20; POINT pt(0,0); BOOL fNC=0; BOOL fWide=1 }
    header = struct.pack('<IiiII', 20, 0, 0, 0, 1)
    files = '\0'.join(os.path.abspath(p) for p in paths) + '\0\0'
    return header + files.encode('utf-16-le')


def copy_result(path):
    """Кладёт файл (CF_HDROP) и его текст (CF_UNICODETEXT) в буфер обмена.
    Возвращает True при успехе. Исключения пробрасываются вызывающему."""
    text = extract_text(path)
    blob = _hdrop_blob([path])
    win32clipboard.OpenClipboard()
    try:
        win32clipboard.EmptyClipboard()
        win32clipboard.SetClipboardData(win32clipboard.CF_HDROP, blob)
        if text:
            win32clipboard.SetClipboardData(win32clipboard.CF_UNICODETEXT, text)
    finally:
        win32clipboard.CloseClipboard()
    return True
