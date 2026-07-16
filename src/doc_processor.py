import os
import re
import itertools
import docx
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.oxml import serialize_part_xml
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from nlp_engine import NLPProcessor, PlaceholderMapper, TAG_PER
from security_utils import (
    MAX_XLSX_CELLS,
    atomic_output,
    check_input_file,
    load_matching_manifest,
    scrub_extended_properties,
    unique_output_path,
    validate_ooxml,
    write_manifest,
)


class DocumentProcessor:
    def __init__(self, nlp_processor: NLPProcessor):
        self.nlp = nlp_processor

    # ------------------------------------------------------------------ #
    #  Анонимизация                                                        #
    # ------------------------------------------------------------------ #
    def process_file(self, file_path, hide_names=True, hide_locations=True,
                     hide_orgs=True, hide_dates=False, smart_contract_mode=False, export_md=False):
        """Анонимизирует .txt/.docx, сохраняет '<имя> [ANON].<ext>' рядом.

        Дефолт smart_contract_mode=False намеренно совпадает с deanonymize_file и
        с тем, что всегда шлёт UI (§10.1): универсальный режим скрывает всё.
        Иначе карта, построенная при восстановлении из оригинала (там режим
        False), не совпала бы с анонимизацией в режиме True → рассинхрон меток.

        export_md=True (только .docx): вместо [ANON].docx создаётся [ANON].md —
        markdown для отправки ИИ (лучше читается моделями). Возвращает путь.
        Восстановление идёт из оригинала (отдельный файл-ключ не создаётся)."""
        file_path = os.path.abspath(file_path)
        file_dir, file_name = os.path.split(file_path)
        name, ext = os.path.splitext(file_name)
        ext = ext.lower()
        if ext == '.doc':
            raise ValueError(
                "Старый формат .doc отключён: его открытие может запускать макросы Word. "
                "Сохраните документ как .docx без макросов."
            )
        supported = {'.txt', '.docx', '.pdf', '.xlsx', '.csv', '.html'}
        if ext not in supported:
            raise ValueError(f"Неподдерживаемое расширение файла: {ext}")
        check_input_file(file_path, text=ext in {'.txt', '.csv', '.html'})
        if ext == '.docx':
            validate_ooxml(file_path, 'docx')
        elif ext == '.xlsx':
            validate_ooxml(file_path, 'xlsx', reject_formulas=True)
        self.nlp.clear_ner_cache()

        # HTML/CSV/PDF intentionally leave the active container and become inert
        # UTF-8 text.  This prevents hidden attributes and spreadsheet formula
        # injection from crossing the anonymisation boundary.
        out_ext = '.md' if export_md and ext == '.docx' else (
            '.txt' if ext in {'.pdf', '.csv', '.html'} else ext
        )
        out_path = unique_output_path(os.path.join(file_dir, f"{name} [ANON]{out_ext}"))
        options = {
            'hide_names': hide_names,
            'hide_locations': hide_locations,
            'hide_orgs': hide_orgs,
            'hide_dates': hide_dates,
            'smart_contract_mode': smart_contract_mode,
            'export_md': export_md and ext == '.docx',
        }
        try:
            with atomic_output(out_path) as temporary:
                if export_md and ext == '.docx':
                    self._process_docx_to_md(file_path, temporary, hide_names,
                                             hide_locations, hide_orgs, hide_dates,
                                             smart_contract_mode)
                elif ext == '.pdf':
                    self._process_pdf(file_path, temporary, hide_names, hide_locations,
                                      hide_orgs, hide_dates, smart_contract_mode)
                elif ext == '.txt':
                    self._process_txt(file_path, temporary, hide_names, hide_locations,
                                      hide_orgs, hide_dates, smart_contract_mode)
                elif ext == '.docx':
                    self._process_docx(file_path, temporary, hide_names, hide_locations,
                                       hide_orgs, hide_dates, smart_contract_mode)
                elif ext == '.xlsx':
                    self._process_xlsx(file_path, temporary, hide_names, hide_locations,
                                       hide_orgs, hide_dates, smart_contract_mode)
                elif ext == '.csv':
                    self._process_csv(file_path, temporary, hide_names, hide_locations,
                                      hide_orgs, hide_dates, smart_contract_mode)
                else:
                    self._process_html(file_path, temporary, hide_names, hide_locations,
                                       hide_orgs, hide_dates, smart_contract_mode)
            write_manifest(out_path, file_path, options)
        except Exception:
            # A document without its sidecar cannot be restored reliably.  Do
            # not leave a half-published result that looks usable.
            if os.path.exists(out_path):
                try:
                    os.remove(out_path)
                except OSError:
                    pass
            raise
        return out_path

    def _process_pdf(self, in_path, out_path, hide_names, hide_locations, hide_orgs, hide_dates, smart_contract_mode):
        """PDF → извлечённый текст → анонимизация → '[ANON].txt'. Скан без
        текстового слоя — понятная ошибка."""
        import pdf_reader
        lines = pdf_reader.extract_lines(in_path)
        if not any(l.strip() for l in lines):
            raise ValueError('В PDF нет текстового слоя (похоже на скан). Нужен PDF '
                             'с распознанным текстом или сначала прогоните OCR.')
        mapper = PlaceholderMapper()
        lines_nl = [l + '\n' for l in lines]
        out_lines = self._anonymize_lines(lines_nl, mapper, hide_names, hide_locations,
                                          hide_orgs, hide_dates, smart_contract_mode)
        with open(out_path, 'w', encoding='utf-8') as f:
            f.writelines(out_lines)
        return mapper

    # ------------------------------------------------------------------ #
    #  Markdown-контур (экспорт для ИИ)                                    #
    # ------------------------------------------------------------------ #
    def _detect_unsupported(self, doc):
        """Содержимое, которое НЕ пройдёт md-контур и грозит утечкой ПДн:
        режим рецензирования, текстовые поля/надписи, поля-конструкторы, а также
        поля/смарт-теги (w:fldSimple/w:smartTag/w:customXml) — их текст
        анонимизатор не видит (paragraph.text = только прямые w:r/w:hyperlink),
        поэтому реальные ПДн из MERGEFIELD утекли бы в ИИ. Проверяем и тело, и
        части сносок/концевых сносок. Пустой список = можно работать."""
        roots = [doc.element]
        for part in self._note_parts(doc):
            roots.append(parse_xml(part.blob))

        def has(tag):
            return any(r.find('.//' + qn(tag)) is not None for r in roots)

        def has_field_with_text(tag):
            for r in roots:
                for el in r.iter(qn(tag)):
                    if any((t.text or '').strip() for t in el.iter(qn('w:t'))):
                        return True
            return False

        issues = []
        if has('w:ins') or has('w:del'):
            issues.append('в документе есть непринятые правки рецензирования — '
                          'примите или отклоните их в Word (Рецензирование → Принять)')
        if has('w:txbxContent'):
            issues.append('в документе есть надписи/текстовые поля — их текст ИИ '
                          'не увидит; перенесите текст из рамок в обычные абзацы')
        if has('w:sdt'):
            issues.append('в документе есть поля-конструкторы (элементы управления) — '
                          'преобразуйте их в обычный текст')
        if (has_field_with_text('w:fldSimple') or has_field_with_text('w:smartTag')
                or has_field_with_text('w:customXml')):
            issues.append('в документе есть автоподставляемые поля или смарт-теги '
                          '(например, данные из шаблона/CRM) — их значения не '
                          'обезличиваются; выделите текст и Ctrl+Shift+F9, чтобы '
                          'превратить поля в обычный текст, затем повторите')
        if has('w:altChunk'):
            issues.append('в документе есть встроенный внешний документ (altChunk) — '
                          'вставьте его содержимое как обычный текст')
        return issues

    def _process_docx_to_md(self, in_path, out_path, hide_names, hide_locations,
                            hide_orgs, hide_dates, smart_contract_mode):
        """Создаёт только [ANON].md. [ANON].docx НЕ создаётся: юрист не должен
        путаться, какой файл отправлять ИИ."""
        from md_serializer import serialize_document

        doc = Document(in_path)
        issues = self._detect_unsupported(doc)
        if issues:
            raise ValueError('Документ нужно подготовить: ' + '; '.join(issues) + '.')

        mapper = PlaceholderMapper()
        self._anonymize_doc(doc, mapper, hide_names, hide_locations, hide_orgs, hide_dates, smart_contract_mode)
        md_text, _blocks, _notes = serialize_document(doc, mapper)
        with open(out_path, 'w', encoding='utf-8') as f:
            f.write(md_text)
        return out_path

    def _build_extra_from_texts(self, texts):
        """Проход-сбор: по всем текстам собирает имена/названия (с метками) и
        строит regex для сквозной протяжки (скрытие пропусков NER по документу)."""
        surfaces, tokens, full = {}, {}, []
        for t in texts:
            if t and t.strip():
                full.append(t)
                s, tk = self.nlp.collect_terms(t)
                for k, v in s.items():
                    surfaces.setdefault(k, v)
                for k, v in tk.items():
                    # Приоритет [ФИО] над [ОРГАНИЗАЦИЯ] при конфликте.
                    if k not in tokens or v == TAG_PER:
                        tokens[k] = v
        return self.nlp.build_extra_patterns(surfaces, tokens, "\n".join(full))

    # Заголовки, включающие/выключающие анонимизацию в smart_contract_mode.
    _SMART_OFF_MARKERS = ("предмет договор",)
    _SMART_ON_MARKERS = (
        "адреса и реквизиты", "реквизиты сторон", "подписи сторон",
        "адреса, реквизиты", "реквизиты банка",
    )

    def _smart_toggle(self, text, is_anonymizing):
        lower_text = text.lower()
        if any(m in lower_text for m in self._SMART_OFF_MARKERS):
            is_anonymizing = False
        if any(m in lower_text for m in self._SMART_ON_MARKERS):
            is_anonymizing = True
        return is_anonymizing

    @staticmethod
    def _read_text(path):
        """Читает текстовый файл терпимо к кодировке: UTF-8 (в т.ч. с BOM), а если
        не вышло — cp1251 (типично, если документ сохранён в старом редакторе).
        Падать на кодировке нельзя — иначе весь пакет обработки прерывается."""
        raw = open(path, 'rb').read()
        try:
            return raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            return raw.decode('cp1251', errors='replace')

    def _read_text_lines(self, path):
        return self._read_text(path).splitlines(keepends=True)

    def _process_txt(self, in_path, out_path, hide_names, hide_locations, hide_orgs, hide_dates, smart_contract_mode):
        lines = self._read_text_lines(in_path)
        mapper = PlaceholderMapper()
        out_lines = self._anonymize_lines(lines, mapper, hide_names, hide_locations, hide_orgs, hide_dates, smart_contract_mode)
        with open(out_path, 'w', encoding='utf-8') as f:
            f.writelines(out_lines)
        return mapper

    def _anonymize_lines(self, lines, mapper, hide_names, hide_locations, hide_orgs, hide_dates, smart_contract_mode):
        extra = self._build_extra_from_texts(lines)
        is_anonymizing = True
        out_lines = []
        for line in lines:
            if smart_contract_mode:
                is_anonymizing = self._smart_toggle(line, is_anonymizing)
            if is_anonymizing and line.strip():
                line = self.nlp.anonymize_text(line, hide_names, hide_locations, hide_orgs,
                                               hide_dates, extra_patterns=extra, mapper=mapper)
            out_lines.append(line)
        return out_lines

    def _iter_table_paragraphs(self, table):
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    yield para
                for nested in cell.tables:
                    yield from self._iter_table_paragraphs(nested)

    def _iter_paragraphs(self, doc):
        """Все параграфы документа: тело, таблицы (в т.ч. вложенные), колонтитулы."""
        for para in doc.paragraphs:
            yield para
        for table in doc.tables:
            yield from self._iter_table_paragraphs(table)
        for section in doc.sections:
            for hf in (
                section.header, section.first_page_header, section.even_page_header,
                section.footer, section.first_page_footer, section.even_page_footer,
            ):
                if hf is None or hf.is_linked_to_previous:
                    continue
                for para in hf.paragraphs:
                    yield para
                for table in hf.tables:
                    yield from self._iter_table_paragraphs(table)

    # ------------------------------------------------------------------ #
    #  Сноски и концевые сноски (footnotes / endnotes)                     #
    # ------------------------------------------------------------------ #
    # python-docx НЕ моделирует эти части: они лежат в отдельных XML-частях
    # (word/footnotes.xml, word/endnotes.xml) и приходят как «сырой» Part с
    # единственным .blob (без .element). Поэтому парсим XML сами, оборачиваем
    # каждый <w:p> в Paragraph (чтобы переиспользовать пофрагментную замену с
    # сохранением форматирования) и сериализуем часть обратно ТЕМ ЖЕ
    # сериализатором, что и весь python-docx (serialize_part_xml) — байты в том
    # же формате, документ остаётся валидным для Word.
    _NOTE_RELTYPES = (RT.FOOTNOTES, RT.ENDNOTES)

    def _note_parts(self, doc):
        """Части сносок/концевых сносок в детерминированном порядке (сноски
        раньше концевых, затем по имени части) — это важно для стабильной
        нумерации меток и, как следствие, точного восстановления из оригинала."""
        rels = [r for r in doc.part.rels.values() if r.reltype in self._NOTE_RELTYPES]

        def rank(rel):
            return (0 if rel.reltype == RT.FOOTNOTES else 1, str(rel.target_part.partname))

        for rel in sorted(rels, key=rank):
            yield rel.target_part

    def _iter_note_texts(self, doc):
        """Текст абзацев сносок/концевых сносок — только для прохода-сбора имён.
        Ничего не меняет."""
        for part in self._note_parts(doc):
            root = parse_xml(part.blob)
            for p_el in root.findall('.//' + qn('w:p')):
                yield Paragraph(p_el, part).text

    def _process_note_parts(self, doc, handle_paragraph):
        """Применяет handle_paragraph(Paragraph) -> bool к каждому абзацу каждой
        note-части. Если хоть один абзац изменён — сериализует часть обратно в
        blob (иначе не трогаем, чтобы документы без сносок остались байт-в-байт).
        Служебные абзацы-разделители (separator/continuationSeparator) без текста
        отсеиваются самим handle_paragraph по .text."""
        for part in self._note_parts(doc):
            root = parse_xml(part.blob)
            changed = False
            for p_el in root.findall('.//' + qn('w:p')):
                if handle_paragraph(Paragraph(p_el, part)):
                    changed = True
            if changed:
                part._blob = serialize_part_xml(root)

    @staticmethod
    def _scrub_core_properties(doc):
        """Чистит метаданные документа (docProps/core.xml): автор, кто изменял,
        заголовок, тема и т.п. Там часто настоящие ФИО/e-mail, но в тело документа
        они не входят, поэтому анонимайзер тела их не видит и они утекли бы в
        [ANON].docx и в буфер обмена при отправке файла ИИ."""
        cp = doc.core_properties
        for attr in ('author', 'last_modified_by', 'title', 'subject',
                     'comments', 'category', 'keywords', 'content_status',
                     'identifier'):
            if getattr(cp, attr, None):
                setattr(cp, attr, '')

    _COMMENT_RELTYPE = RT.COMMENTS

    def _comment_parts(self, doc):
        """XML-части рецензий (word/comments.xml). python-docx их не моделирует."""
        for rel in doc.part.rels.values():
            if rel.reltype == self._COMMENT_RELTYPE:
                yield rel.target_part

    def _anonymize_comment_parts(self, doc, hide_names, hide_locations, hide_orgs, hide_dates, mapper):
        """Анонимизирует текст комментариев Word и стирает имя/инициалы автора.
        Без этого рецензии («Согласовано с Ивановым, ИНН …») и ФИО рецензента
        уходили в [ANON].docx нетронутыми — критическая утечка."""
        for part in self._comment_parts(doc):
            root = parse_xml(part.blob)
            changed = False
            for cmt in root.iter(qn('w:comment')):
                for attr in ('w:author', 'w:initials'):
                    if cmt.get(qn(attr)):
                        cmt.set(qn(attr), '')
                        changed = True
            p_elems = root.findall('.//' + qn('w:p'))
            extra = self._build_extra_from_texts([Paragraph(p, part).text for p in p_elems])
            for p_el in p_elems:
                if self._anonymize_paragraph(Paragraph(p_el, part), hide_names,
                                             hide_locations, hide_orgs, hide_dates, extra, mapper):
                    changed = True
            if changed:
                part._blob = serialize_part_xml(root)

    def _process_docx(self, in_path, out_path, hide_names, hide_locations, hide_orgs, hide_dates, smart_contract_mode):
        doc = Document(in_path)
        # Тот же барьер, что и в md-контуре: поля (w:fldSimple), надписи
        # (w:txbxContent), смарт-теги, поля-конструкторы и непринятые правки
        # анонимизатор НЕ видит (paragraph.text = только w:r|w:hyperlink), поэтому
        # их ПДн ушли бы в [ANON].docx НЕобезличенными. Раньше барьер стоял только
        # на md-пути — обычный .docx-режим (он по умолчанию) тихо протекал.
        issues = self._detect_unsupported(doc)
        if issues:
            raise ValueError('Документ нужно подготовить: ' + '; '.join(issues) + '.')
        mapper = PlaceholderMapper()
        self._anonymize_doc(doc, mapper, hide_names, hide_locations, hide_orgs, hide_dates, smart_contract_mode)
        self._anonymize_comment_parts(doc, hide_names, hide_locations, hide_orgs, hide_dates, mapper)
        self._scrub_core_properties(doc)
        doc.save(out_path)
        scrub_extended_properties(out_path)
        return mapper

    def _anonymize_doc(self, doc, mapper, hide_names, hide_locations, hide_orgs, hide_dates, smart_contract_mode):
        """Анонимизирует объект Document на месте, наполняя mapper. Порядок
        обхода детерминирован — важно для повторного построения карты из оригинала."""
        # Проход 1: собрать имена/названия по ВСЕМУ документу — тело, таблицы,
        # колонтитулы И сноски. Сквозная протяжка ловит пропуски NER и падежи, в
        # т.ч. когда имя есть в теле, но пропущено в сноске (или наоборот).
        extra = self._build_extra_from_texts(
            itertools.chain((p.text for p in self._iter_paragraphs(doc)),
                            self._iter_note_texts(doc)))

        # Тело.
        is_anonymizing = True
        for para in doc.paragraphs:
            if smart_contract_mode:
                is_anonymizing = self._smart_toggle(para.text, is_anonymizing)
            if is_anonymizing:
                self._anonymize_paragraph(para, hide_names, hide_locations, hide_orgs, hide_dates, extra, mapper)

        # Таблицы (fail-safe: по умолчанию анонимизируем).
        for table in doc.tables:
            self._process_table(table, hide_names, hide_locations, hide_orgs, hide_dates,
                                smart_contract_mode, True, extra, mapper)

        # Колонтитулы (всегда, вне smart-режима).
        for section in doc.sections:
            for hf in (
                section.header, section.first_page_header, section.even_page_header,
                section.footer, section.first_page_footer, section.even_page_footer,
            ):
                if hf is None or hf.is_linked_to_previous:
                    continue
                for para in hf.paragraphs:
                    self._anonymize_paragraph(para, hide_names, hide_locations, hide_orgs, hide_dates, extra, mapper)
                for table in hf.tables:
                    self._process_table(table, hide_names, hide_locations, hide_orgs, hide_dates,
                                        False, True, extra, mapper)

        # Сноски и концевые сноски (всегда, вне smart-режима): это служебные
        # примечания, а не «тело договора», но ПДн в них утекают так же.
        self._process_note_parts(doc, lambda para: self._anonymize_paragraph(
            para, hide_names, hide_locations, hide_orgs, hide_dates, extra, mapper))

    def _process_table(self, table, hide_names, hide_locations, hide_orgs, hide_dates,
                       smart_contract_mode, is_anonymizing, extra, mapper):
        # Объединённые ячейки (gridSpan/vMerge) python-docx отдаёт одним и тем же
        # объектом в нескольких позициях row.cells. Без дедупликации такая ячейка
        # анонимизировалась бы дважды: второй проход прогонял бы NER уже по МЕТКАМ
        # ([ФИО_1]) и портил их. Обрабатываем каждый физический <w:tc> один раз.
        seen = set()
        for row in table.rows:
            for cell in row.cells:
                if id(cell._tc) in seen:
                    continue
                seen.add(id(cell._tc))
                for para in cell.paragraphs:
                    if smart_contract_mode:
                        is_anonymizing = self._smart_toggle(para.text, is_anonymizing)
                    if is_anonymizing:
                        self._anonymize_paragraph(para, hide_names, hide_locations, hide_orgs, hide_dates, extra, mapper)
                for nested in cell.tables:
                    is_anonymizing = self._process_table(
                        nested, hide_names, hide_locations, hide_orgs, hide_dates,
                        smart_contract_mode, is_anonymizing, extra, mapper)
        return is_anonymizing

    def _anonymize_paragraph(self, paragraph, hide_names, hide_locations, hide_orgs, hide_dates, extra, mapper):
        """Анонимизирует абзац на месте. Возвращает True, если что-то заменено
        (нужно для сносок: сериализуем часть обратно только при изменениях)."""
        if not paragraph.text.strip():
            return False
        replacements = self.nlp.extract_entities(
            paragraph.text, hide_names, hide_locations, hide_orgs, hide_dates=hide_dates, extra_patterns=extra)
        if not replacements:
            return False
        # Метки-номера присваиваем слева направо (replacements идут справа-налево).
        for rep in reversed(replacements):
            rep['ins'] = mapper.placeholder(rep['category'], rep['text'])
        self._apply_span_replacements(paragraph, replacements)
        return True

    def _apply_span_replacements(self, paragraph, replacements):
        """Применяет замены к runs параграфа с сохранением форматирования.
        replacements — отсортированы справа-налево, каждая с ключом 'ins'
        (строка-вставка). Используется и для анонимизации, и для восстановления."""
        all_runs = []
        try:
            for item in paragraph.iter_inner_content():
                if isinstance(item, docx.text.run.Run):
                    all_runs.append(item)
                elif hasattr(item, 'runs'):
                    all_runs.extend(item.runs)
        except AttributeError:
            all_runs = paragraph.runs

        run_intervals = []
        current_idx = 0
        for r in all_runs:
            length = len(r.text)
            run_intervals.append({'run': r, 'start': current_idx, 'stop': current_idx + length, 'text': r.text})
            current_idx += length

        for rep in replacements:
            start_idx, stop_idx, ins = rep['start'], rep['stop'], rep['ins']
            overlapping = [ri for ri in run_intervals
                           if max(start_idx, ri['start']) < min(stop_idx, ri['stop'])]
            if not overlapping:
                continue
            for i, ri in enumerate(overlapping):
                local_start = max(0, start_idx - ri['start'])
                local_stop = min(len(ri['text']), stop_idx - ri['start'])
                if i == 0:
                    new_text = ri['text'][:local_start] + ins + ri['text'][local_stop:]
                else:
                    new_text = ri['text'][:local_start] + ri['text'][local_stop:]
                ri['text'] = new_text
                self._set_run_text_preserving(ri['run'], new_text)

    @staticmethod
    def _set_run_text_preserving(run, new_text):
        """Как `run.text = new_text`, но НЕ уничтожает не-текстовые элементы рана:
        встроенные картинки/подписи (w:drawing/w:pict/w:object) и разрывы страниц
        (w:br type=page/column). Штатный сеттер python-docx вызывает clear_content()
        и стирает их вместе с текстом — терялись подпись-картинка и пагинация, причём
        деанонимизация их уже не возвращала. Текст-только раны (99% случаев) идут
        прежним быстрым путём."""
        r = run._r
        preserved = False
        for child in r:
            if child.tag in (qn('w:drawing'), qn('w:pict'), qn('w:object')):
                preserved = True
                break
            if child.tag == qn('w:br') and child.get(qn('w:type')) in ('page', 'column'):
                preserved = True
                break
        if not preserved:
            run.text = new_text
            return
        # Есть сохраняемые элементы: правим только текстовые узлы w:t. Весь новый
        # текст кладём в первый w:t (сохраняя его позицию относительно картинки),
        # остальные w:t обнуляем.
        t_elems = r.findall(qn('w:t'))
        if not t_elems:
            return
        t_elems[0].text = new_text
        t_elems[0].set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        for extra in t_elems[1:]:
            extra.text = ''

    # ------------------------------------------------------------------ #
    #  Деанонимизация (восстановление)                                    #
    # ------------------------------------------------------------------ #
    def deanonymize_file(self, new_path, source_path, hide_names=True, hide_locations=True,
                         hide_orgs=True, hide_dates=False, smart_contract_mode=False):
        """Восстанавливает данные в 'new_path' (документ от ИИ с метками).
        source_path — исходный документ, привязанный паспортом *.umbra.json.
        Сохраняет '<имя> [DEANON].<ext>'. Возвращает (out_path, stats), где
        stats = {'restored': int, 'unresolved': int, 'samples': [...]}."""
        new_path = os.path.abspath(new_path)
        source_path = os.path.abspath(source_path)
        source_ext = os.path.splitext(source_path)[1].lower()
        answer_ext = os.path.splitext(new_path)[1].lower()
        if source_ext not in {'.txt', '.docx', '.pdf', '.xlsx', '.csv', '.html'}:
            raise ValueError("Неподдерживаемый оригинал для восстановления.")
        if answer_ext not in {'.txt', '.md', '.docx', '.xlsx'}:
            raise ValueError(
                "Ответы .doc/.html/.csv запрещены из-за активного содержимого. "
                "Сохраните ответ как .txt, .md, безопасный .docx или .xlsx."
            )
        check_input_file(new_path, text=answer_ext in {'.txt', '.md'})
        check_input_file(source_path, text=source_ext in {'.txt', '.csv', '.html'})
        if source_ext == '.docx':
            validate_ooxml(source_path, 'docx')
        elif source_ext == '.xlsx':
            validate_ooxml(source_path, 'xlsx', reject_formulas=True)
        if answer_ext == '.docx':
            validate_ooxml(new_path, 'docx')
        elif answer_ext == '.xlsx':
            validate_ooxml(new_path, 'xlsx', reject_formulas=True)

        manifest = load_matching_manifest(source_path)
        options = manifest['options']
        hide_names = bool(options.get('hide_names', True))
        hide_locations = bool(options.get('hide_locations', True))
        hide_orgs = bool(options.get('hide_orgs', True))
        hide_dates = bool(options.get('hide_dates', False))
        smart_contract_mode = bool(options.get('smart_contract_mode', False))
        self.nlp.clear_ner_cache()

        mapping = self._load_mapping(source_path, hide_names, hide_locations, hide_orgs, hide_dates, smart_contract_mode)
        if not mapping:
            raise ValueError("Карта замен пуста — нечего восстанавливать.")

        file_dir, file_name = os.path.split(new_path)
        name, ext = os.path.splitext(file_name)
        # Убираем возможные суффиксы [ANON]/[ОТВЕТ], чтобы имя было аккуратным.
        clean = name.replace(' [ANON]', '').replace(' [ОТВЕТ]', '')
        out_ext = '.txt' if ext.lower() == '.md' and source_ext != '.docx' else ext
        out_path = unique_output_path(os.path.join(file_dir, f"{clean} [DEANON]{out_ext}"))

        tol_re = self._build_tolerant_re(mapping)
        if ext.lower() == '.md':
            # Ответ ИИ в markdown: полный контур с восстановлением .docx
            # (форматирование оригинала + правки ИИ). Возвращает СЕССИЮ —
            # сохранение происходит после экрана верификации в GUI.
            if os.path.splitext(source_path)[1].lower() != '.docx':
                # Плоский путь для .txt-источника: метки → значения прямо в md.
                with atomic_output(out_path) as temporary:
                    stats = self._deanon_txt(new_path, temporary, mapping, tol_re)
                return out_path, stats
            return self._deanon_md(new_path, source_path, mapping,
                                   hide_names, hide_locations, hide_orgs, hide_dates, smart_contract_mode)
        with atomic_output(out_path) as temporary:
            if ext.lower() == '.txt':
                stats = self._deanon_txt(new_path, temporary, mapping, tol_re)
            elif ext.lower() == '.docx':
                stats = self._deanon_docx(new_path, temporary, mapping, tol_re)
            else:
                stats = self._deanon_xlsx(new_path, temporary, mapping, tol_re)
        return out_path, stats

    # ------------------------------------------------------------------ #
    #  Markdown-контур (восстановление .docx из ответа ИИ)                #
    # ------------------------------------------------------------------ #
    def _deanon_md(self, ai_md_path, source_path, mapping,
                   hide_names, hide_locations, hide_orgs, hide_dates, smart_contract_mode):
        """Трёхстороннее слияние: правки из md от ИИ вживляются в копию
        оригинала, затем метки заменяются реальными данными В ПАМЯТИ.
        Возвращает MdDeanonSession (needs_confirmation=True): документ ещё НЕ
        сохранён — GUI показывает отчёт и вызывает session.save()."""
        from md_serializer import serialize_document, LABEL_RE_CI
        import md_merge

        # 1. Копия оригинала + карта ОДНИМ вызовом _anonymize_doc: карта, база
        #    слияния и базовые блоки рождаются вместе — рассинхрон исключён.
        doc = Document(source_path)
        mapper = PlaceholderMapper()
        self._anonymize_doc(doc, mapper, hide_names, hide_locations, hide_orgs, hide_dates, smart_contract_mode)
        self._anonymize_comment_parts(doc, hide_names, hide_locations, hide_orgs,
                                      hide_dates, mapper)
        self._scrub_core_properties(doc)

        # 2. Карта восстановления — из повторной анонимизации оригинала. Файл-ключ
        #    .umbra убран (один выходной файл проще); поэтому восстановление
        #    рассчитывает, что оригинал не меняли после обезличивания.
        mapping = mapper.mapping
        warnings = []

        # 3. База слияния — повторная сериализация (детерминирована).
        _base_md, base_blocks, notes = serialize_document(doc, mapper)

        # UTF-8 (в т.ч. с BOM) — норма; cp1251 — если юрист сохранил ответ в
        # старом редакторе. Падать на кодировке нельзя.
        raw = open(ai_md_path, 'rb').read()
        try:
            ai_text = raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            ai_text = raw.decode('cp1251', errors='replace')

        # 4. Метки ответа ⊆ карте? Ловит «выбрали не тот оригинал» и выдуманные
        #    ИИ метки ДО каких-либо правок документа.
        ai_labels = {'[' + m.group(1).upper() + '_' + m.group(2) + ']'
                     for m in LABEL_RE_CI.finditer(ai_text)}
        known_cats = {k.split('_')[0].lstrip('[') for k in mapping}
        unknown_labels = sorted(l for l in ai_labels
                                if l not in mapping and l.split('_')[0].lstrip('[') in known_cats)

        # 5. Слияние (удаления откладываются до подтверждения).
        report, deletions = md_merge.merge_ai_md(
            ai_text, base_blocks, mapping, notes, self._apply_span_replacements)

        # 6. Скан галлюцинаций: в правленых/вставленных фрагментах не должно
        #    быть живых ПДн — только метки. NER/regex по изменённым текстам.
        #    Сущности, попадающие ВНУТРЬ меток (regex ловит цифры из «[ФИО_12]»),
        #    не считаются: сверяем интервалы с координатами меток.
        suspects = []
        for text in report.get('changed_texts', []):
            label_spans = [m.span() for m in LABEL_RE_CI.finditer(text)]
            ents = self.nlp.extract_entities(text)
            for e in ents:
                if any(s <= e['start'] and e['stop'] <= t for s, t in label_spans):
                    continue
                val = e['text'].strip()
                if '[' in val or ']' in val:
                    continue
                if any(val in v or v in val for v in mapping.values()):
                    continue
                if val not in suspects:
                    suspects.append(val)

        # 7. Деанонимизация в памяти.
        tol_re = self._build_tolerant_re(mapping)
        restored, unresolved = 0, []
        if tol_re is not None:
            for para in self._iter_paragraphs(doc):
                r, u = self._deanonymize_paragraph(para, mapping, tol_re)
                restored += r
                unresolved += u

            # Сноски: деанонимизируем ТО ЖЕ дерево notes.by_id, куда md_merge внёс
            # правки ИИ, и помечаем часть изменённой — notes.commit() в save()
            # запишет именно его. Повторный парс part.blob здесь был бы затёрт
            # этим commit'ом, и метки остались бы в итоговом .docx (критбаг).
            if notes is not None:
                for (kind, nid), paras in notes.by_id.items():
                    changed = False
                    for para in paras:
                        r, u = self._deanonymize_paragraph(para, mapping, tol_re)
                        restored += r
                        unresolved += u
                        if r > 0:
                            changed = True
                    if changed:
                        notes.mark_changed(kind, nid)
            c_restored, c_unresolved = self._deanonymize_comment_parts(doc, mapping, tol_re)
            restored += c_restored
            unresolved += c_unresolved

        # 8. Остаточный скан: метки, пережившие деанон (сильные искажения ИИ),
        #    включая англоязычные подделки (NAME_1) без скобок. Обходим тело И
        #    сноски (иначе метка в сноске молча ушла бы в итоговый документ).
        residual = set()
        scan_texts = [p.text for p in self._iter_paragraphs(doc)]
        if notes is not None:
            for paras in notes.by_id.values():
                scan_texts.extend(p.text for p in paras)
        for part in self._comment_parts(doc):
            root = parse_xml(part.blob)
            scan_texts.extend(
                Paragraph(p_el, part).text
                for p_el in root.findall('.//' + qn('w:p'))
            )
        for text in scan_texts:
            for m in LABEL_RE_CI.finditer(text):
                residual.add(m.group(0))
            for m in re.finditer(r'\b(NAME|ORG|COMPANY|AMOUNT|SUM|PHONE|EMAIL|ADDRESS)[ _-]?\d+\b',
                                 text, re.IGNORECASE):
                residual.add(m.group(0))

        stats = {'restored': restored, 'unresolved': len(unresolved),
                 'samples': unresolved[:5]}

        # Краткая сводка предупреждений для статус-строки (модального экрана
        # верификации больше нет — решение владельца «проще»). Слияние безопасно
        # по построению (инвариант меток, безопасная склейка), поэтому документ не
        # искажён; здесь лишь ЧЕСТНО сообщаем о том, что стоит перепроверить.
        notices = []
        if report.get('truncation_suspect'):
            notices.append('возможно, ИИ вернул документ не целиком (покрытие %.0f%%) — '
                           'недостающее взято из оригинала' % (100 * report.get('coverage', 0)))
        if unknown_labels:
            notices.append('ИИ использовал неизвестные метки (%s) — данных для них нет'
                           % ', '.join(unknown_labels[:5]))
        resid = sorted(residual)
        if resid:
            notices.append('метки не распознаны и остались в тексте: %s' % ', '.join(resid[:5]))
        if suspects:
            notices.append('возможно, ИИ вписал вымышленные данные: %s — проверьте'
                           % ', '.join(suspects[:5]))
        if report.get('label_mismatch'):
            notices.append('блоков с изменёнными ИИ метками: %d — оставлены из оригинала '
                           '(данные не искажены)' % report['label_mismatch'])
        manual_del = sum(1 for d in deletions if not d.get('auto'))
        if manual_del:
            notices.append('ИИ предложил удалить блоков: %d — требуется подтверждение' % manual_del)
        for c in report.get('table_conflicts', []):
            notices.append('таблица: ' + c)
        notices.extend(report.get('warnings', []))
        notices.extend(warnings)
        stats['notices'] = notices

        file_dir, file_name = os.path.split(ai_md_path)
        clean = os.path.splitext(file_name)[0].replace(' [ANON]', '').replace(' [ОТВЕТ]', '')
        out_path = unique_output_path(os.path.join(file_dir, f"{clean} [DEANON].docx"))

        return MdDeanonSession(doc, notes, out_path, stats, deletions)


    def _build_tolerant_re(self, mapping):
        """Строит regex, устойчивый к «причёсыванию» меток ИИ: разный регистр и
        разделитель ([ФИО_1], [ФИО 1], [ФИО-1], [фио1]). Категории берём из карты."""
        cats = set()
        for key in mapping:
            m = re.match(r'\[(.+)_\d+\]$', key)
            if m:
                cats.add(m.group(1))
        if not cats:
            return None
        # Длинные категории первыми, чтобы не отхватить префикс (ИНДЕКС до ИНН).
        alt = '|'.join(re.escape(c) for c in sorted(cats, key=len, reverse=True))
        return re.compile(r'\[\s*(' + alt + r')[\s_\-]*(\d+)\s*\]', re.IGNORECASE)

    def _load_mapping(self, source_path, hide_names, hide_locations, hide_orgs, hide_dates, smart_contract_mode):
        ext = os.path.splitext(source_path)[1].lower()
        # The manifest has already verified the exact original and restored the
        # original options.  Re-running the deterministic anonymiser now builds
        # the same in-memory mapping without persisting personal data.
        mapper = PlaceholderMapper()
        if ext == '.docx':
            doc = Document(source_path)
            self._anonymize_doc(doc, mapper, hide_names, hide_locations, hide_orgs, hide_dates, smart_contract_mode)
            self._anonymize_comment_parts(doc, hide_names, hide_locations, hide_orgs,
                                          hide_dates, mapper)
        elif ext == '.txt':
            lines = self._read_text_lines(source_path)
            self._anonymize_lines(lines, mapper, hide_names, hide_locations, hide_orgs, hide_dates, smart_contract_mode)
        elif ext == '.pdf':
            import pdf_reader
            lines = [l + '\n' for l in pdf_reader.extract_lines(source_path)]
            self._anonymize_lines(lines, mapper, hide_names, hide_locations, hide_orgs, hide_dates, smart_contract_mode)
        elif ext in ('.xlsx', '.csv', '.html'):
            # Эти форматы обезличиваются функциями, которые СРАЗУ пишут файл, а нам
            # нужна только карта. Пишем во временный файл и удаляем его: карта
            # детерминирована, поэтому совпадает с той, что была при анонимизации.
            # Без этой ветки GUI находил оригинал .xlsx/.csv/.html, а восстановление
            # падало с «Источник восстановления: …» — формат был односторонним.
            import tempfile
            handlers = {'.xlsx': self._process_xlsx, '.csv': self._process_csv,
                        '.html': self._process_html}
            fd, tmp_out = tempfile.mkstemp(suffix=ext)
            os.close(fd)
            try:
                mapper = handlers[ext](source_path, tmp_out, hide_names, hide_locations,
                                       hide_orgs, hide_dates, smart_contract_mode)
            finally:
                if os.path.exists(tmp_out):
                    os.remove(tmp_out)
        else:
            raise ValueError("Источник восстановления: .docx/.txt/.pdf/.xlsx/.csv/.html.")
        return mapper.mapping

    @staticmethod
    def _canonical(match):
        """Приводит найденную (возможно искажённую) метку к каноничному виду
        [КАТЕГОРИЯ_N] для поиска в карте."""
        return '[' + match.group(1).upper() + '_' + match.group(2) + ']'

    def _deanon_docx(self, in_path, out_path, mapping, tol_re):
        doc = Document(in_path)
        restored, unresolved = 0, []
        if tol_re is not None:
            for para in self._iter_paragraphs(doc):
                r, u = self._deanonymize_paragraph(para, mapping, tol_re)
                restored += r
                unresolved += u

            # Сноски и концевые сноски (отдельные XML-части, с записью обратно).
            acc = {'restored': 0, 'unresolved': []}

            def handle(para):
                r, u = self._deanonymize_paragraph(para, mapping, tol_re)
                acc['restored'] += r
                acc['unresolved'].extend(u)
                return r > 0   # менять blob есть смысл только если что-то восстановлено

            self._process_note_parts(doc, handle)
            restored += acc['restored']
            unresolved += acc['unresolved']
            c_restored, c_unresolved = self._deanonymize_comment_parts(doc, mapping, tol_re)
            restored += c_restored
            unresolved += c_unresolved
        self._scrub_core_properties(doc)
        doc.save(out_path)
        scrub_extended_properties(out_path)
        return {'restored': restored, 'unresolved': len(unresolved), 'samples': unresolved[:5]}

    def _deanonymize_comment_parts(self, doc, mapping, tol_re):
        restored, unresolved = 0, []
        for part in self._comment_parts(doc):
            root = parse_xml(part.blob)
            changed = False
            for comment in root.iter(qn('w:comment')):
                for attr in ('w:author', 'w:initials'):
                    if comment.get(qn(attr)):
                        comment.set(qn(attr), '')
                        changed = True
            for p_el in root.findall('.//' + qn('w:p')):
                r, u = self._deanonymize_paragraph(
                    Paragraph(p_el, part), mapping, tol_re)
                restored += r
                unresolved += u
                changed = changed or r > 0
            if changed:
                part._blob = serialize_part_xml(root)
        return restored, unresolved

    def _deanonymize_paragraph(self, paragraph, mapping, tol_re):
        text = paragraph.text
        if '[' not in text:
            return 0, []
        reps, restored, unresolved = [], 0, []
        for m in tol_re.finditer(text):
            val = mapping.get(self._canonical(m))
            if val is not None:
                reps.append({'start': m.start(), 'stop': m.end(), 'ins': val})
                restored += 1
            else:
                unresolved.append(m.group(0))
        if reps:
            reps.sort(key=lambda x: x['start'], reverse=True)
            self._apply_span_replacements(paragraph, reps)
        return restored, unresolved

    def _restore_text_multipass(self, text, mapping, tol_re, stats):
        RESTORE_MAX_PASSES = 10
        def repl(m):
            val = mapping.get(self._canonical(m))
            if val is not None:
                stats['restored'] += 1
                return val
            stats['unresolved'] += 1
            if len(stats['samples']) < 5:
                stats['samples'].append(m.group(0))
            return m.group(0)

        if tol_re is not None:
            for _ in range(RESTORE_MAX_PASSES):
                prev = text
                text = tol_re.sub(repl, text)
                if text == prev:
                    break
        return text

    def _deanon_txt(self, in_path, out_path, mapping, tol_re):
        text = self._read_text(in_path)
        stats = {'restored': 0, 'unresolved': 0, 'samples': []}
        text = self._restore_text_multipass(text, mapping, tol_re, stats)
        with open(out_path, 'w', encoding='utf-8') as f:
            f.write(text)
        return stats

    def _deanon_xlsx(self, in_path, out_path, mapping, tol_re):
        import openpyxl
        stats = {'restored': 0, 'unresolved': 0, 'samples': []}
        wb = openpyxl.load_workbook(str(in_path))
        for ws in wb.worksheets:
            if ws.max_row * ws.max_column > MAX_XLSX_CELLS:
                raise ValueError(f'Лист «{ws.title}» превышает лимит в {MAX_XLSX_CELLS} ячеек.')
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value and isinstance(cell.value, str):
                        cell.value = self._restore_text_multipass(cell.value, mapping, tol_re, stats)
                    if cell.comment is not None:
                        cell.comment.text = self._restore_text_multipass(
                            cell.comment.text or '', mapping, tol_re, stats)
                        cell.comment.author = ''
            for hf in (ws.oddHeader, ws.oddFooter, ws.evenHeader, ws.evenFooter,
                       ws.firstHeader, ws.firstFooter):
                for part in (hf.left, hf.center, hf.right):
                    if getattr(part, 'text', None):
                        part.text = self._restore_text_multipass(part.text, mapping, tol_re, stats)
        self._scrub_workbook_properties(wb)
        wb.save(str(out_path))
        scrub_extended_properties(out_path)
        return stats

    def _process_html(self, in_path, out_path, hide_names, hide_locations, hide_orgs, hide_dates, smart_contract_mode):
        from html.parser import HTMLParser
        mapper = PlaceholderMapper()

        class VisibleText(HTMLParser):
            BLOCKS = {'address', 'article', 'aside', 'blockquote', 'br', 'div',
                      'footer', 'h1', 'h2', 'h3', 'h4', 'h5', 'h6', 'header',
                      'li', 'main', 'p', 'section', 'table', 'tr'}
            SKIP = {'script', 'style', 'template', 'noscript'}

            def __init__(self):
                super().__init__(convert_charrefs=True)
                self.parts = []
                self.skip_depth = 0

            def handle_starttag(self, tag, attrs):
                tag = tag.lower()
                if tag in self.SKIP:
                    self.skip_depth += 1
                elif not self.skip_depth and tag in self.BLOCKS:
                    self.parts.append('\n')

            def handle_endtag(self, tag):
                tag = tag.lower()
                if tag in self.SKIP and self.skip_depth:
                    self.skip_depth -= 1
                elif not self.skip_depth and tag in self.BLOCKS:
                    self.parts.append('\n')

            def handle_data(self, data):
                if not self.skip_depth:
                    self.parts.append(data)

        src = self._read_text(in_path)
        parser = VisibleText()
        parser.feed(src)
        visible = ''.join(parser.parts)
        lines = visible.splitlines(keepends=True)
        output = self._anonymize_lines(lines, mapper, hide_names, hide_locations,
                                       hide_orgs, hide_dates, smart_contract_mode)
        with open(out_path, 'w', encoding="utf-8") as f:
            f.writelines(output)
        return mapper

    def _process_csv(self, in_path, out_path, hide_names, hide_locations, hide_orgs, hide_dates, smart_contract_mode):
        import csv
        import io
        mapper = PlaceholderMapper()
        lines = []
        cell_count = 0
        source = self._read_text(in_path)
        try:
            dialect = csv.Sniffer().sniff(source[:65_536], delimiters=',;\t|')
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(io.StringIO(source, newline=''), dialect)
        for row in reader:
            cell_count += len(row)
            if cell_count > MAX_XLSX_CELLS:
                raise ValueError(f'CSV превышает лимит в {MAX_XLSX_CELLS} ячеек.')
            clean = [cell.replace('\r', ' ').replace('\n', ' ').replace('\t', ' ')
                     for cell in row]
            lines.append('\t'.join(clean) + '\n')
        output = self._anonymize_lines(lines, mapper, hide_names, hide_locations,
                                       hide_orgs, hide_dates, smart_contract_mode)
        with open(out_path, "w", encoding="utf-8", newline="\n") as f:
            f.writelines(output)
        return mapper

    def _process_xlsx(self, in_path, out_path, hide_names, hide_locations, hide_orgs, hide_dates, smart_contract_mode):
        import openpyxl
        mapper = PlaceholderMapper()

        def anon(s):
            return self.nlp.anonymize_text(s, hide_names, hide_locations, hide_orgs,
                                           hide_dates, mapper=mapper)

        wb = openpyxl.load_workbook(str(in_path))
        for sheet_number, ws in enumerate(wb.worksheets, start=1):
            if ws.max_row * ws.max_column > MAX_XLSX_CELLS:
                raise ValueError(f'Лист «{ws.title}» превышает лимит в {MAX_XLSX_CELLS} ячеек.')
            # Sheet titles are copied into hidden extended properties and may
            # themselves contain names/customer identifiers.  Use inert generic
            # names; the original workbook remains the source of truth.
            ws.title = f'Лист {sheet_number}'
            for row in ws.iter_rows():
                for cell in row:
                    v = cell.value
                    if isinstance(v, str):
                        cell.value = anon(v)
                    elif isinstance(v, (int, float)) and not isinstance(v, bool):
                        # Числовые ПДн (ИНН, телефон, счёт, ОГРН, СНИЛС) Excel хранит
                        # как ЧИСЛО — прежний строковый гейт их пропускал. Скрываем
                        # длинные числа (>=10 цифр — идентификаторы); короткие
                        # (цены, количества, годы) не трогаем, чтобы не портить данные.
                        s = str(v)
                        if sum(ch.isdigit() for ch in s) >= 10:
                            a = anon(s)
                            if a != s:
                                cell.value = a
                    # Примечание к ячейке: текст и автор — тоже утечка.
                    cmt = cell.comment
                    if cmt is not None:
                        if cmt.text:
                            cmt.text = anon(cmt.text)
                        if getattr(cmt, 'author', None):
                            cmt.author = anon(cmt.author)
            self._anonymize_xlsx_headers_footers(ws, anon)
        self._scrub_workbook_properties(wb)
        wb.save(str(out_path))
        scrub_extended_properties(out_path)
        return mapper

    @staticmethod
    def _scrub_workbook_properties(wb):
        props = wb.properties
        for attr in ('creator', 'lastModifiedBy', 'title', 'subject', 'description',
                     'keywords', 'category', 'identifier', 'contentStatus'):
            if hasattr(props, attr):
                setattr(props, attr, '')
        wb.defined_names.clear()

    @staticmethod
    def _anonymize_xlsx_headers_footers(ws, anon):
        """Колонтитулы листа (печатные верх/низ) тоже несут ПДн — обезличиваем."""
        for hf in (ws.oddHeader, ws.oddFooter, ws.evenHeader, ws.evenFooter,
                   ws.firstHeader, ws.firstFooter):
            for part in (hf.left, hf.center, hf.right):
                if getattr(part, 'text', None):
                    part.text = anon(part.text)


class MdDeanonSession:
    """Результат md-восстановления: документ собран в памяти, метки уже заменены
    реальными данными. GUI по флагу needs_confirmation вызывает save() и
    показывает краткий итог (stats['notices']) — без модального экрана верификации
    (решение владельца: проще). Корректность обеспечивает само слияние (инвариант
    сохранности меток, безопасная склейка), а не ручная вычитка."""

    needs_confirmation = True     # GUI по этому флагу вызывает save()

    def __init__(self, doc, notes, out_path, stats, deletions):
        self.doc = doc
        self.notes = notes
        self.out_path = out_path
        self.stats = stats
        self.deletions = deletions          # [{'bi','preview','ref','auto'}]

    def save(self, confirmed_deletions=None):
        """Применяет автоудаления (склейка абзацев — безопасны по построению:
        второй абзац удаляется, только если его метки перенесены в первый),
        включает обновление полей Word и сохраняет [DEANON].docx. Возвращает путь."""
        import md_merge
        md_merge.apply_deletions(self.deletions, confirmed_deletions)
        if self.notes is not None:
            self.notes.commit()
        try:
            settings = self.doc.settings.element
            if settings.find(qn('w:updateFields')) is None:
                upd = parse_xml(
                    '<w:updateFields xmlns:w="http://schemas.openxmlformats.org/'
                    'wordprocessingml/2006/main" w:val="true"/>')
                settings.append(upd)
        except Exception:
            pass
        with atomic_output(self.out_path) as temporary:
            self.doc.save(temporary)
            scrub_extended_properties(temporary)
        return self.out_path

