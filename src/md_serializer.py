"""Экспорт УЖЕ анонимизированного Document в Markdown ([ANON].md).

Принцип: md — транспорт для ИИ, а не формат хранения. Сериализация ДЕТЕРМИНИРОВАНА:
при восстановлении оригинал повторно анонимизируется и сериализуется — получаются
те же блоки, что были отправлены (база для трёхстороннего слияния в md_merge).

Возвращает (md_text, blocks, notes): blocks — список блоков с привязкой к живым
объектам документа, notes — реестр сносок (для записи правок обратно в XML-части).

Блок: {'kind': 'para'|'table'|'frozen'|'footnote',
       'md': str,      # блок в [ANON].md (без завершающих переводов строк)
       'norm': str,    # нормализованный текст для выравнивания (см. normalize_for_align)
       'ref':  Paragraph | Table | (kind, note_id),
       'meta': {...}}  # para: num_literal, heading; table: dims

Защита меток: [ФИО_1] в md-синтаксисе выглядит как ссылка, а «[ФИО_1]: …» в начале
строки — как reference definition. Поэтому при экранировании метки временно
заменяются на PUA-символы (U+E000+i) и не могут пострадать; после метки перед
':' или '(' вставляется word-joiner U+2060 — чужой md-рендер (по пути к ИИ) не
съест строку как definition/ссылку.
"""

import re
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table

MD_FORMAT_VERSION = 1

# Широкий образец метки (категория кириллицей/латиницей + номер) — для защиты и
# нормализации. Намеренно шире карты: ловит и искажённые ИИ варианты.
# Категория может содержать ВНУТРЕННИЕ подчёркивания ([ИНН_ЮЛ], [КОД_ПОДР],
# [МЕСТО_РОЖДЕНИЯ]…), но начинается и заканчивается буквой — иначе такие метки были
# бы невидимы инварианту целостности меток при слиянии, и реальный ИНН мог потеряться.
LABEL_RE = re.compile(r'\[\s*([A-ZА-ЯЁ][A-ZА-ЯЁa-zа-яё_]{0,18}[A-ZА-ЯЁa-zа-яё])[\s_\-]*(\d{1,4})\s*\]')
# Регистронезависимый вариант — ТОЛЬКО для нормализации при выравнивании и для
# защиты с фильтром по известным меткам (иначе «[статья 5]» в тексте пострадала бы).
LABEL_RE_CI = re.compile(r'\[\s*([A-Za-zА-Яа-яЁё][A-Za-zА-Яа-яЁё_]{0,18}[A-Za-zА-Яа-яЁё])[\s_\-]*(\d{1,4})\s*\]')

IMG_TOKEN_RE = re.compile(r'\[\[РИСУНОК_(\d+)\]\]')
# Ссылка на сноску в md ([^f12]); в тексте docx её нет (это XML-элемент), поэтому
# при выравнивании и при записи текста обратно её нужно игнорировать/срезать.
FOOTREF_RE = re.compile(r'\[\^[fe]?\d{1,4}\]')
_PUA_BASE = 0xE000
# word-joiner (U+2060): невидим, разрывает md-синтаксис «]:», «](». Строим через
# chr() — в исходнике нет невидимых символов (правило проекта).
_WJ = chr(0x2060)
_PUA_BEFORE_PUNCT_RE = re.compile('([' + chr(_PUA_BASE) + '-' + chr(0xF8FF) + '])([:(])')

# Ведущая нумерация пункта: «7.», «7.2.1.», «а)», «1)» (+ пробел).
_ENUM_PREFIX_RE = re.compile(r'^\s*(?:\d{1,3}(?:\.\d{1,3})*[.)]|[а-яa-z][.)])\s+', re.IGNORECASE)

_TYPO_MAP = str.maketrans({
    '«': '"', '»': '"', '„': '"', '“': '"', '”': '"',
    '’': "'", '‘': "'",
    '–': '-', '—': '-', '−': '-',
    # NBSP / узкий NBSP / тонкий пробел / таб — через chr(): в исходнике нет
    # невидимых символов (правило проекта).
    chr(0xA0): ' ', chr(0x202F): ' ', chr(0x2009): ' ', chr(9): ' ',
    # ё=е: модели сплошь пишут «е» вместо «ё» — это не правка текста.
    'ё': 'е', 'Ё': 'Е',
    # Нулевой ширины (word-joiner U+2060 мы сами добавляем в _finish_md_text,
    # плюс U+200B/U+FEFF от рендеров ИИ) — удаляем (None), иначе неизменённый
    # абзац «[МЕТКА]:» с невидимым символом ошибочно считается правкой.
    chr(0x2060): None, chr(0x200B): None, chr(0xFEFF): None,
})


# --------------------------------------------------------------------------- #
#  Общие помощники (используются и md_merge)                                   #
# --------------------------------------------------------------------------- #

def protect_labels(text, registry=None, known=None):
    """Заменяет метки на PUA-символы. registry: {pua_char: канон '[КАТ_N]'} —
    передавайте один словарь на документ, одинаковые метки получают один символ.

    known — множество канонических меток из карты замен. Если задано, защищаются
    (и потом канонизируются при восстановлении) ТОЛЬКО они: обычный текст в
    скобках («[Статья 5]») не пострадает. Без known защищается всё похожее на
    метку (режим сериализации собственного анонимизированного текста)."""
    if registry is None:
        registry = {}
    rev = {v: k for k, v in registry.items()}

    def sub(m):
        canon = '[' + m.group(1).upper() + '_' + m.group(2) + ']'
        if known is not None and canon not in known:
            return m.group(0)
        ch = rev.get(canon)
        if ch is None:
            ch = chr(_PUA_BASE + len(registry))
            registry[ch] = canon
            rev[canon] = ch
        return ch

    pattern = LABEL_RE_CI if known is not None else LABEL_RE
    return pattern.sub(sub, text), registry


def restore_labels(text, registry):
    for ch, canon in registry.items():
        text = text.replace(ch, canon)
    return text


def normalize_for_align(text):
    """Нормализация для выравнивания блоков: канонические метки, унификация
    типографики (ИИ часто «чинит» кавычки/тире — это не правка), без ведущей
    нумерации, без токенов рисунков, схлопнутые пробелы."""
    t = LABEL_RE_CI.sub(lambda m: '[' + m.group(1).upper() + '_' + m.group(2) + ']', text)
    t = IMG_TOKEN_RE.sub(' ', t)
    t = FOOTREF_RE.sub(' ', t)
    t = t.translate(_TYPO_MAP)
    t = _ENUM_PREFIX_RE.sub('', t)
    t = re.sub(r'\s+', ' ', t).strip()
    return t


def _escape_md(text):
    """Экранирует спецсимволы md в обычном тексте (метки уже под PUA-защитой)."""
    text = text.replace('\\', '\\\\')
    text = re.sub(r'([*_\[\]<>|#`~])', r'\\\1', text)
    return text


def _escape_line_start(line):
    """Гасит блочный синтаксис в начале строки: списки, цитаты, заголовки,
    нумерация «7.» (иначе CommonMark перенумерует пункты договора)."""
    m = re.match(r'^(\s*)(\d{1,3}(?:\.\d{1,3})*)([.)])(\s)', line)
    if m:
        return m.group(1) + m.group(2).replace('.', '\\.') + '\\' + m.group(3) + m.group(4) + line[m.end():]
    m = re.match(r'^(\s*)([-+>])(\s)', line)
    if m:
        return m.group(1) + '\\' + m.group(2) + m.group(3) + line[m.end():]
    return line


def _finish_md_text(text):
    """Финализация md-строки: восстановление меток произойдёт снаружи; здесь —
    word-joiner после PUA-метки перед ':'/'(' (см. шапку модуля)."""
    return _PUA_BEFORE_PUNCT_RE.sub(lambda m: m.group(1) + _WJ + m.group(2), text)


def _safe_int(v):
    """int(v) или None вместо исключения — для XML-атрибутов, которые в кривых
    (часто конвертированных не из Word) документах бывают пустыми/нечисловыми.
    Без этого один битый w:val роняет весь экспорт md."""
    try:
        return int(v)
    except (TypeError, ValueError):
        return None


# --------------------------------------------------------------------------- #
#  Нумерация (numbering.xml)                                                   #
# --------------------------------------------------------------------------- #

_FMT_ALPHA_RU = 'абвгдежзиклмнопрстуфхцчшщэюя'   # без ё,й,ъ,ы,ь — как в Word
_FMT_ALPHA_EN = 'abcdefghijklmnopqrstuvwxyz'
_ROMAN = ((1000, 'M'), (900, 'CM'), (500, 'D'), (400, 'CD'), (100, 'C'), (90, 'XC'),
          (50, 'L'), (40, 'XL'), (10, 'X'), (9, 'IX'), (5, 'V'), (4, 'IV'), (1, 'I'))


def _to_roman(n):
    out = []
    for v, s in _ROMAN:
        while n >= v:
            out.append(s)
            n -= v
    return ''.join(out) or 'I'


def _fmt_number(n, fmt):
    if fmt == 'decimal' or fmt == 'decimalZero':
        return str(n)
    if fmt == 'lowerLetter':
        return _FMT_ALPHA_EN[(n - 1) % 26]
    if fmt == 'upperLetter':
        return _FMT_ALPHA_EN[(n - 1) % 26].upper()
    if fmt == 'russianLower':
        return _FMT_ALPHA_RU[(n - 1) % len(_FMT_ALPHA_RU)]
    if fmt == 'russianUpper':
        return _FMT_ALPHA_RU[(n - 1) % len(_FMT_ALPHA_RU)].upper()
    if fmt == 'lowerRoman':
        return _to_roman(n).lower()
    if fmt == 'upperRoman':
        return _to_roman(n)
    return None   # неизвестный формат — консервативно без номера


class NumberingResolver:
    """Вычисляет ВИДИМЫЕ номера пунктов по numbering.xml.

    Консервативен: при isLgl/lvlRestart/numStyleLink/неизвестном формате номер
    не подставляется, а numId попадает в self.report (неверный номер в md хуже
    отсутствующего — ИИ начнёт ссылаться на несуществующие пункты)."""

    def __init__(self, doc):
        self.report = []
        self._counters = {}       # numId -> {ilvl: счётчик}
        self._levels = {}         # numId -> {ilvl: {'fmt','text','start'}}
        self._bad = set()         # numId, для которых номер не считаем
        try:
            self._load(doc)
        except Exception as e:
            self.report.append(f'нумерация не разобрана ({e}); номера пунктов в md не подставлены')
            self._levels = {}

    def _load(self, doc):
        num_root = None
        for rel in doc.part.rels.values():
            if rel.reltype == RT.NUMBERING:
                part = rel.target_part
                num_root = getattr(part, 'element', None)
                if num_root is None:
                    num_root = parse_xml(part.blob)
                break
        if num_root is None:
            return

        abstract = {}   # abstractNumId -> {ilvl: lvl-info} | None (=подозрителен)
        for an in num_root.findall(qn('w:abstractNum')):
            aid = an.get(qn('w:abstractNumId'))
            if an.find(qn('w:numStyleLink')) is not None:
                abstract[aid] = None
                continue
            levels, suspicious = {}, False
            for lvl in an.findall(qn('w:lvl')):
                ilvl = int(lvl.get(qn('w:ilvl')))
                if lvl.find(qn('w:isLgl')) is not None or lvl.find(qn('w:lvlRestart')) is not None:
                    suspicious = True
                fmt_el = lvl.find(qn('w:numFmt'))
                txt_el = lvl.find(qn('w:lvlText'))
                start_el = lvl.find(qn('w:start'))
                levels[ilvl] = {
                    'fmt': fmt_el.get(qn('w:val')) if fmt_el is not None else 'decimal',
                    'text': txt_el.get(qn('w:val')) if txt_el is not None else '',
                    'start': int(start_el.get(qn('w:val'))) if start_el is not None else 1,
                }
            abstract[aid] = None if suspicious else levels

        for num in num_root.findall(qn('w:num')):
            num_id = num.get(qn('w:numId'))
            aid_el = num.find(qn('w:abstractNumId'))
            if aid_el is None:
                continue
            levels = abstract.get(aid_el.get(qn('w:val')))
            if levels is None:
                self._bad.add(num_id)
                self.report.append(f'сложная нумерация (numId={num_id}) — номера не подставлены')
                continue
            levels = {k: dict(v) for k, v in levels.items()}
            for ov in num.findall(qn('w:lvlOverride')):
                ilvl = int(ov.get(qn('w:ilvl')))
                so = ov.find(qn('w:startOverride'))
                if so is not None and ilvl in levels:
                    levels[ilvl]['start'] = int(so.get(qn('w:val')))
                if ov.find(qn('w:lvl')) is not None:
                    # Полная переопределённая lvl — редкость; консервативно.
                    self._bad.add(num_id)
            if num_id in self._bad:
                self.report.append(f'переопределение уровня (numId={num_id}) — номера не подставлены')
                continue
            self._levels[num_id] = levels

    @staticmethod
    def _num_pr(paragraph):
        """(numId, ilvl) абзаца: сначала прямой numPr, затем numPr стиля."""
        try:
            pPr = paragraph._p.pPr
            if pPr is not None and pPr.numPr is not None:
                num_id = pPr.numPr.numId
                ilvl = pPr.numPr.ilvl
                if num_id is not None:
                    lv = _safe_int(ilvl.val) if ilvl is not None else 0
                    return str(num_id.val), (lv if lv is not None else 0)
        except Exception:
            pass
        try:
            style = paragraph.style
            while style is not None:
                s_pPr = style.element.find(qn('w:pPr'))
                if s_pPr is not None:
                    numPr = s_pPr.find(qn('w:numPr'))
                    if numPr is not None:
                        nid = numPr.find(qn('w:numId'))
                        ilv = numPr.find(qn('w:ilvl'))
                        if nid is not None:
                            return (nid.get(qn('w:val')),
                                    int(ilv.get(qn('w:val'))) if ilv is not None else 0)
                style = style.base_style
        except Exception:
            pass
        return None, None

    def resolve(self, paragraph):
        """Видимый номер абзаца ('7.2.1.', 'а)', '-') или None."""
        num_id, ilvl = self._num_pr(paragraph)
        if num_id is None or num_id == '0' or num_id in self._bad:
            return None
        levels = self._levels.get(num_id)
        if levels is None or ilvl not in levels:
            return None

        counters = self._counters.setdefault(num_id, {})
        counters[ilvl] = counters.get(ilvl, levels[ilvl]['start'] - 1) + 1
        for deeper in list(counters):
            if deeper > ilvl:
                del counters[deeper]

        info = levels[ilvl]
        if info['fmt'] == 'bullet':
            return '-'
        text = info['text']

        def sub_placeholder(m):
            lv = int(m.group(1)) - 1
            lv_info = levels.get(lv)
            if lv_info is None:
                return m.group(0)
            val = counters.get(lv, lv_info['start'])
            s = _fmt_number(val, lv_info['fmt'])
            return s if s is not None else m.group(0)

        rendered = re.sub(r'%(\d)', sub_placeholder, text)
        if '%' in rendered or not rendered.strip():
            return None
        return rendered


# --------------------------------------------------------------------------- #
#  Абзац -> markdown                                                           #
# --------------------------------------------------------------------------- #

def _heading_level(paragraph):
    """Уровень заголовка 1..6 или None: outlineLvl абзаца/стиля, имя стиля."""
    pPr = paragraph._p.pPr
    if pPr is not None:
        el = pPr.find(qn('w:outlineLvl'))
        if el is not None:
            lvl = _safe_int(el.get(qn('w:val')))
            return lvl + 1 if lvl is not None and 0 <= lvl <= 5 else None
    try:
        style = paragraph.style
        while style is not None:
            s_pPr = style.element.find(qn('w:pPr'))
            if s_pPr is not None:
                el = s_pPr.find(qn('w:outlineLvl'))
                if el is not None:
                    lvl = _safe_int(el.get(qn('w:val')))
                    return lvl + 1 if lvl is not None and 0 <= lvl <= 5 else None
            m = re.match(r'(?:heading|заголовок)\s*(\d)', style.name or '', re.IGNORECASE)
            if m:
                lvl = int(m.group(1))
                return lvl if 1 <= lvl <= 6 else None
            style = style.base_style
    except Exception:
        pass
    return None


def _run_flags(run):
    return (bool(run.bold), bool(run.italic),
            bool(run.font.strike), bool(run.underline))


_MARKS = (('**', '**'), ('*', '*'), ('~~', '~~'), ('<u>', '</u>'))


def _wrap_flags(text, flags):
    if not text.strip():
        return text
    lead = text[:len(text) - len(text.lstrip())]
    trail = text[len(text.rstrip()):]
    core = text.strip()
    for i, on in enumerate(flags):
        if on:
            core = _MARKS[i][0] + core + _MARKS[i][1]
    return lead + core + trail


class _ParaRenderer:
    """Инлайновый рендер абзаца: runs -> md (жирный/курсив/зачёркнутый/подчёркнутый),
    рисунки -> [[РИСУНОК_n]], ссылки на сноски -> [^fN]/[^eN]."""

    def __init__(self, registry, counters):
        self.registry = registry
        self.counters = counters   # {'img': int} — сквозной счётчик рисунков

    def _render_run(self, run_el, paragraph, pieces):
        """Токены одного run в порядке следования (текст с начертанием, рисунки,
        сноски, таб/перенос)."""
        from docx.text.run import Run
        flags = _run_flags(Run(run_el, paragraph))
        for c in run_el:
            ct = c.tag
            if ct == qn('w:t'):
                pieces.append((c.text or '', flags))
            elif ct == qn('w:tab') or ct == qn('w:ptab'):
                pieces.append(('\t', flags))
            elif ct == qn('w:noBreakHyphen'):
                # paragraph.text даёт обычный «-»; чтобы render не расходился с
                # доменом анонимизатора (иначе неразрывный дефис ломает round-trip:
                # неизменённый абзац примут за правку и вырежут дефис), эмитим тот
                # же символ как часть слова.
                pieces.append(('-', flags))
            elif ct == qn('w:cr'):
                pieces.append(('\n', None))
            elif ct == qn('w:br'):
                # paragraph.text даёт '\n' только для textWrapping (или без типа);
                # разрыв страницы/колонки в текст НЕ попадает — иначе render
                # разошёлся бы с paragraph.text.
                br_type = c.get(qn('w:type'))
                if br_type in (None, 'textWrapping'):
                    pieces.append(('\n', None))
            elif ct == qn('w:drawing') or ct == qn('w:pict'):
                self.counters['img'] += 1
                pieces.append(('[[РИСУНОК_%d]]' % self.counters['img'], None))
            elif ct == qn('w:footnoteReference'):
                pieces.append(('[^f' + c.get(qn('w:id')) + ']', None))
            elif ct == qn('w:endnoteReference'):
                pieces.append(('[^e' + c.get(qn('w:id')) + ']', None))

    def render(self, paragraph):
        # КРИТИЧНО (утечка ПДн): обходим ТОЛЬКО прямые w:r и runs внутри прямых
        # w:hyperlink — ровно то, что видит анонимизатор (paragraph.text в
        # python-docx = xpath('w:r | w:hyperlink')). Текст внутри w:fldSimple
        # (кэш MERGEFIELD), w:smartTag, w:customXml НЕ анонимизируется, поэтому в
        # md его выводить нельзя — иначе сырые ПДн ушли бы в ИИ. Такие поля ловит
        # _detect_unsupported и блокирует экспорт; здесь — второй рубеж (не эмитим).
        pieces = []     # (text, flags|None)  None = сырой токен (не экранировать)
        for child in paragraph._p:
            if child.tag == qn('w:r'):
                self._render_run(child, paragraph, pieces)
            elif child.tag == qn('w:hyperlink'):
                for r in child.findall(qn('w:r')):
                    self._render_run(r, paragraph, pieces)

        # Склейка соседних кусков с одинаковыми флагами.
        merged = []
        for text, flags in pieces:
            if merged and merged[-1][1] == flags and flags is not None:
                merged[-1][0] += text
            else:
                merged.append([text, flags])

        out = []
        for text, flags in merged:
            if flags is None:
                out.append(text)
                continue
            prot, _ = protect_labels(text, self.registry)
            esc = _escape_md(prot)
            out.append(_wrap_flags(esc, flags))
        line = ''.join(out)
        # Мягкие переводы строк (w:br) -> двойной пробел + \n (md hard break).
        line = line.replace('\n', '  \n')
        return line


# --------------------------------------------------------------------------- #
#  Таблицы                                                                     #
# --------------------------------------------------------------------------- #

def _table_is_simple(table):
    """Таблица «простая» (можно в GFM): без gridSpan/vMerge/вложенных таблиц."""
    tbl = table._tbl
    if tbl.findall('.//' + qn('w:tbl')):
        return False
    for tc_pr in tbl.findall('.//' + qn('w:tcPr')):
        gs = tc_pr.find(qn('w:gridSpan'))
        if gs is not None and int(gs.get(qn('w:val'), '1')) > 1:
            return False
        if tc_pr.find(qn('w:vMerge')) is not None:
            return False
    return True


def _cell_text_md(cell, renderer):
    parts = [renderer.render(p) for p in cell.paragraphs]
    txt = '<br>'.join(p for p in parts if p.strip()) or ' '
    return txt.replace('|', '\\|').replace('  \n', '<br>')


def _table_to_gfm(table, renderer):
    rows = []
    for row in table.rows:
        cells = [_cell_text_md(c, renderer) for c in row.cells]
        rows.append('| ' + ' | '.join(cells) + ' |')
    if not rows:
        return ''
    ncols = len(table.rows[0].cells)
    delim = '| ' + ' | '.join(['---'] * ncols) + ' |'
    return rows[0] + '\n' + delim + '\n' + '\n'.join(rows[1:]) if len(rows) > 1 else rows[0] + '\n' + delim


def _html_escape(s):
    """Экранирование для HTML-контекста (в отличие от _escape_md — для markdown)."""
    return s.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')


def _frozen_cell_content(cell):
    """Содержимое ячейки замороженной таблицы для HTML: анонимизированный текст
    абзацев (paragraph.text — тот же безопасный домен w:r|w:hyperlink, что видит
    анонимизатор) + РЕКУРСИВНО вложенные таблицы, HTML-экранированный. Начертание
    в замороженных таблицах опускаем: они помечены «структуру не менять» и при
    импорте не правятся, важны только содержимое и метки."""
    parts = [p.text for p in cell.paragraphs if p.text.strip()]
    txt = _html_escape(' '.join(parts))
    for nested in cell.tables:
        txt += _table_to_frozen_html(nested, None)
    return txt or ' '


def _table_to_frozen_html(table, renderer=None):
    """Сложная таблица «замораживается» HTML-блоком: ИИ видит содержимое (включая
    вложенные таблицы), но предупреждён не менять структуру. При импорте такая
    таблица не правится поячеечно (берётся из оригинала), поэтому renderer не
    нужен — текст берём напрямую и HTML-экранируем (раньше сюда попадало
    markdown-экранирование, ломавшее HTML)."""
    out = ['<table data-umbra="структуру не менять">']
    for row in table.rows:
        out.append('<tr>')
        seen = set()
        for cell in row.cells:
            if id(cell._tc) in seen:
                continue
            seen.add(id(cell._tc))
            out.append('<td>' + _frozen_cell_content(cell) + '</td>')
        out.append('</tr>')
    out.append('</table>')
    return '\n'.join(out)


def _table_norm(table):
    texts = []
    for row in table.rows:
        for cell in row.cells:
            texts.append(cell.text)
    return normalize_for_align(' '.join(texts))


def _table_dims(table):
    return (len(table.rows), len(table.columns))


# --------------------------------------------------------------------------- #
#  Сноски                                                                      #
# --------------------------------------------------------------------------- #

_NOTE_RELTYPES = {RT.FOOTNOTES: 'f', RT.ENDNOTES: 'e'}


class NotesRegistry:
    """Части сносок с распарсенными XML-деревьями. Держит Paragraph-обёртки
    живыми до конца слияния; commit() сериализует изменённые части обратно."""

    def __init__(self, doc):
        self.parts = []            # [(part, root, kind)]
        self.by_id = {}            # ('f'|'e', id) -> [Paragraph]
        rels = [r for r in doc.part.rels.values() if r.reltype in _NOTE_RELTYPES]
        rels.sort(key=lambda r: (0 if r.reltype == RT.FOOTNOTES else 1,
                                 str(r.target_part.partname)))
        for rel in rels:
            kind = _NOTE_RELTYPES[rel.reltype]
            part = rel.target_part
            root = parse_xml(part.blob)
            self.parts.append([part, root, kind, False])
            tag = qn('w:footnote') if kind == 'f' else qn('w:endnote')
            for note in root.findall(tag):
                if note.get(qn('w:type')):     # separator/continuation — служебные
                    continue
                nid = note.get(qn('w:id'))
                paras = [Paragraph(p, part) for p in note.findall('.//' + qn('w:p'))]
                if paras:
                    self.by_id[(kind, nid)] = paras

    def mark_changed(self, kind, nid):
        for entry in self.parts:
            if entry[2] == kind:
                entry[3] = True

    def commit(self):
        from docx.opc.oxml import serialize_part_xml
        for part, root, kind, changed in self.parts:
            if changed:
                part._blob = serialize_part_xml(root)


# --------------------------------------------------------------------------- #
#  Главная точка входа                                                         #
# --------------------------------------------------------------------------- #

_AI_HEADER = """\
> **Инструкция для ИИ-ассистента.** Это юридический документ, в котором
> персональные данные заменены метками вида `[ФИО_1]`, `[ОРГАНИЗАЦИЯ_2]`,
> `[СУММА_1]`. Правила работы:
> 1. НИКОГДА не изменяйте, не переводите и не удаляйте метки — возвращайте их
>    в точности как есть, включая номер.
> 2. Не меняйте структуру таблиц с пометкой «структуру не менять».
> 3. Служебные обозначения `[[РИСУНОК_N]]` и сноски `[^fN]` оставляйте на месте.
> 4. Верните документ ЦЕЛИКОМ, не сокращайте не изменённые разделы.
"""


def _label_legend(mapper):
    if mapper is None or not getattr(mapper, 'mapping', None):
        return ''
    labels = sorted(mapper.mapping.keys())
    return ('> В документе использованы метки: ' + ', '.join('`%s`' % l for l in labels) + '\n')


def serialize_document(doc, mapper=None, include_header=True):
    """Сериализует АНОНИМИЗИРОВАННЫЙ Document в markdown.

    Возвращает (md_text, blocks, notes_registry). Детерминирована: одинаковый
    документ -> одинаковые блоки (важно для повторной базы при импорте)."""
    registry = {}                       # PUA -> канон метки
    counters = {'img': 0}
    renderer = _ParaRenderer(registry, counters)
    resolver = NumberingResolver(doc)
    notes = NotesRegistry(doc)

    blocks = []
    body = doc.element.body
    for child in body.iterchildren():
        if child.tag == qn('w:p'):
            para = Paragraph(child, doc)
            raw_text = para.text
            line = renderer.render(para)
            if not line.strip():
                continue
            heading = _heading_level(para)
            num_literal = None
            if heading is None:
                num = resolver.resolve(para)
                if num == '-':
                    num_literal = '-'
                    line = '\\- ' + line
                elif num is not None:
                    num_literal = num
                    # Литеральный номер с экранированными «.»/«)» — CommonMark
                    # не примет его за элемент списка и не перенумерует.
                    esc = re.sub(r'([.)])', r'\\\1', num)
                    line = esc + ' ' + line
                else:
                    line = _escape_line_start(line)
                md = line
            else:
                md = '#' * heading + ' ' + line
            blocks.append({
                'kind': 'para',
                'md': _finish_md_text(md),
                'norm': normalize_for_align(raw_text),
                'ref': para,
                'meta': {'heading': heading, 'num_literal': num_literal},
            })
        elif child.tag == qn('w:tbl'):
            table = Table(child, doc)
            simple = _table_is_simple(table)
            if simple:
                md = _table_to_gfm(table, renderer)
                kind = 'table'
            else:
                md = _table_to_frozen_html(table, renderer)
                kind = 'frozen'
            blocks.append({
                'kind': kind,
                'md': _finish_md_text(md),
                'norm': _table_norm(table),
                'ref': table,
                'meta': {'dims': _table_dims(table)},
            })

    # Сноски/концевые сноски: определения в конце документа.
    note_blocks = []
    for (kind, nid), paras in sorted(notes.by_id.items()):
        text_md = ' '.join(renderer.render(p) for p in paras if p.text.strip())
        if not text_md.strip():
            continue
        raw = ' '.join(p.text for p in paras)
        md = '[^%s%s]: %s' % (kind, nid, text_md)
        note_blocks.append({
            'kind': 'footnote',
            'md': _finish_md_text(md),
            'norm': normalize_for_align(raw),
            'ref': (kind, nid),
            'meta': {},
        })

    # Сборка текста: блоки разделяются пустой строкой.
    parts = []
    if include_header:
        parts.append(_AI_HEADER + _label_legend(mapper))
    parts.extend(b['md'] for b in blocks)
    if note_blocks:
        parts.append('---')
        parts.extend(b['md'] for b in note_blocks)
    md_text = '\n\n'.join(parts) + '\n'
    md_text = restore_labels(md_text, registry)

    blocks.extend(note_blocks)
    for b in blocks:
        b['md'] = restore_labels(b['md'], registry)

    return md_text, blocks, notes
